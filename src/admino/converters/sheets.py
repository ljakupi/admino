"""XLSX and CSV converters (GH-188): one Markdown table per worksheet, with caps.

Also home of the table rules every converter shares (``table_text``: XLSX
sheets, CSV files, DOCX tables) and of ``format_cell``.

Inputs: the path of a stored ``xlsx`` or ``csv`` upload, the part writer and the
conversion options (unused: sheets have no pages or labels).
Outputs: one text part (none when there is no content); returns no page count.
XLSX: ``## <sheet name>`` and its table (or ``(empty sheet)``) per worksheet in
workbook order, chartsheets skipped. CSV: one table, no heading. A table whose
kept cells hold nothing within the column cap is its notes alone.

Security notes:
- The OOXML zip-bomb guard runs before openpyxl opens a workbook; openpyxl
  reads it read-only with cached values (no formula is evaluated).
- Caps read at call time from ``common``: ``MAX_SHEETS``, ``MAX_TABLE_ROWS``,
  ``MAX_TABLE_COLUMNS``, ``MAX_CELL_CHARS`` (via ``clean_cell``); a cut is
  stated in a note. A table stops being built (and stops reading rows) as
  soon as the characters of its kept cells pass ``MAX_TEXT_CHARS``, and a
  workbook once its sections do (``text_too_large``), so shared strings
  repeated across many cells can't inflate memory beyond about the cap.
- CPU traps: the CSV dialect is sniffed from the first 8 KiB only
  (``csv.Sniffer``'s regexes are quadratic), and at most
  ``MAX_SHEET_ROW_INDEX`` rows (Excel's maximum) are read per worksheet:
  openpyxl yields an empty row for every missing index, so a single
  ``<row r="1000000000000">`` would otherwise loop until the timeout.
- Per-row and per-cell work is bounded by the caps, not by what a few bytes
  of XML claim. XLSX rows are read ``MAX_TABLE_COLUMNS + 1`` columns wide:
  openpyxl otherwise pads every row out to its last cell's column (up to
  18,278), so ``<row><c r="ZZZ1"/></row>`` would cost thousands of cells.
  A row whose values are all None or ``""`` is skipped before any cell is
  formatted. Within one table, a ``str`` longer than ``MAX_CELL_CHARS`` is
  cleaned once per distinct value (an LRU of ``_LONG_VALUE_MEMO_SIZE``
  values; ``str`` caches its hash), so one huge shared string referenced by
  every cell is cleaned once, not once per cell.
- Trade-off of the XLSX width: the columns note fires when a kept row has
  content in column ``MAX_TABLE_COLUMNS + 1`` (51), or when the worksheet's
  declared ``<dimension>`` ends right of it (column 52 or later, whatever the
  row: the dimension covers the whole sheet). Content further right is never
  read; in a sheet without a dimension (or with a rows-only one) it is not
  noted either when column 51 is empty.
- Apart from that column check, the declared dimension is dropped before any
  row is read, so a lying ``<dimension>`` can't cut real rows off or add
  row or cell work (a huge one only adds the note).
- Any openpyxl, XML, ZIP or CSV error (while loading or while reading rows)
  is ``corrupted_file``; the library's message is dropped (``from None``).
"""

from __future__ import annotations

import csv
import datetime as dt
import functools
import itertools
from typing import TYPE_CHECKING, Final

import openpyxl

from admino.converters import common
from admino.converters.common import ConversionError, clean_cell, markdown_table
from admino.converters.ooxml import check_archive

if TYPE_CHECKING:
    from collections.abc import Iterable, Sequence
    from pathlib import Path
    from typing import BinaryIO

    from admino.converters.common import ConversionOptions, PartWriter

# Excel's last row, and the most rows read per sheet: openpyxl yields one empty
# row for every index a sheet skips, so a hostile index would loop for hours.
MAX_SHEET_ROW_INDEX: Final = 1_048_576
_CSV_SAMPLE_CHARS: Final = 8 * 1024
_CSV_DELIMITERS: Final = ",;\t|"
_EMPTY_SHEET: Final = "(empty sheet)"
# Distinct long values whose cleaned form one table remembers: enough for a
# sheet's repeated shared strings, bounded whatever the file holds.
_LONG_VALUE_MEMO_SIZE: Final = 1024
_SECTION_SEPARATOR: Final = "\n\n"
# Below this magnitude a float's integral value prints exactly as an int.
_EXACT_INTEGRAL_FLOAT: Final = 1e15


def format_cell(value: object) -> str:
    """Format one cell value as text (before ``clean_cell``).

    None is empty; bools are ``TRUE``/``FALSE``; integral floats under 1e15
    print as ints, other floats as ``repr``; a datetime at midnight is its
    date, any other datetime is ISO with a space; dates and times are ISO;
    everything else (timedelta and str included) is ``str()``.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < _EXACT_INTEGRAL_FLOAT:
            return str(int(value))
        return repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return str(value)


def table_text(rows: Iterable[Sequence[object]], *, columns_cut: bool = False) -> str:
    """Build a capped Markdown table from raw rows, with its notes; ``""`` when empty.

    A row whose values are all None or ``""`` is skipped unformatted. In the
    other rows a None is an empty cell and every other value is formatted
    (``format_cell``) and cleaned (``clean_cell``), a ``str`` longer than
    ``MAX_CELL_CHARS`` once per distinct value in this call (at most
    ``_LONG_VALUE_MEMO_SIZE`` remembered); rows whose cells are all empty
    are dropped; the first ``MAX_TABLE_ROWS`` non-empty rows are
    kept, each cut to ``MAX_TABLE_COLUMNS`` cells, then the table is trimmed
    to its rightmost non-empty column. Notes follow after a blank line, one
    per line: the rows note when another non-empty row follows the kept ones,
    the columns note when a kept row has content beyond the column cap, or
    when ``columns_cut`` is set (the caller knows of content beyond it that
    the rows don't show: an XLSX sheet's declared dimension). When no kept
    cell within the column cap holds content, the text is the notes alone
    (``""`` without notes).

    Raises:
        ConversionError: ``text_too_large`` as soon as the kept cells hold more
            than ``MAX_TEXT_CHARS`` characters (no further row is read).
    """
    max_rows = common.MAX_TABLE_ROWS
    max_columns = common.MAX_TABLE_COLUMNS
    max_chars = common.MAX_TEXT_CHARS
    max_cell_chars = common.MAX_CELL_CHARS
    # Per call: a long value's cleaning costs its length (a 9M-character
    # shared string about 11 ms), and one value can fill every cell.
    clean_long = functools.lru_cache(maxsize=_LONG_VALUE_MEMO_SIZE)(_cell_text)
    kept: list[list[str]] = []
    chars = 0
    rows_cut = False
    for row in rows:
        # Both counts run in C. None first: padded rows are all None, and
        # comparing a None with "" is the slow comparison.
        nones = row.count(None)
        if nones == len(row) or nones + row.count("") == len(row):
            continue
        # A None is "" unformatted: a sparse row costs its values, not its width.
        # (Formatting the 50 Nones of a one-cell row cost about 6 µs a row.)
        cells = [
            ""
            if value is None
            else clean_long(value)
            if isinstance(value, str) and len(value) > max_cell_chars
            else _cell_text(value)
            for value in row
        ]
        if not any(cells):
            continue
        if len(kept) == max_rows:
            # Stop at the first non-empty row past the cap: the rest is never read.
            rows_cut = True
            break
        columns_cut = columns_cut or any(cells[max_columns:])
        kept.append(cells[:max_columns])
        chars += sum(len(cell) for cell in kept[-1])
        if chars > max_chars:
            raise ConversionError("text_too_large")
    notes = []
    if rows_cut:
        notes.append(f"[Only the first {max_rows} rows are included.]")
    if columns_cut:
        notes.append(f"[Only the first {max_columns} columns are included.]")
    width = max((_used_width(row) for row in kept), default=0)
    if not width:
        return "\n".join(notes)
    table = markdown_table([row[:width] for row in kept])
    return table + "\n\n" + "\n".join(notes) if notes else table


def _cell_text(value: object) -> str:
    """One raw value as a table cell: formatted, then cleaned."""
    return clean_cell(format_cell(value))


def _used_width(cells: list[str]) -> int:
    """The number of cells up to and including the last non-empty one."""
    return max((index + 1 for index, cell in enumerate(cells) if cell), default=0)


def convert_xlsx(path: Path, writer: PartWriter, options: ConversionOptions) -> int | None:
    """Write the workbook's worksheets as Markdown tables in one text part.

    Returns:
        None: workbooks have no pages.

    Raises:
        ConversionError: ``archive_too_large`` or ``password_protected``
            (zip-bomb guard), ``corrupted_file`` (unreadable workbook or
            sheet), ``text_too_large`` (more than ``MAX_TEXT_CHARS``
            characters).
    """
    check_archive(path)
    # openpyxl refuses a path without an xlsx extension (stored uploads have
    # none), so it gets the open file; the file is closed whatever happens.
    with path.open("rb") as file:
        try:
            text = _workbook_text(file)
        except (ConversionError, MemoryError):
            raise
        except Exception:
            # openpyxl's error surface isn't closed (zipfile, XML parser,
            # KeyError for a missing part, value errors), and read-only mode
            # parses each sheet lazily while its rows are read.
            raise ConversionError("corrupted_file") from None
    writer.add_text(text)
    return None


def _workbook_text(file: BinaryIO) -> str:
    """The sections of every worksheet (at most ``MAX_SHEETS``) joined by a blank line."""
    workbook = openpyxl.load_workbook(file, read_only=True, data_only=True)
    try:
        worksheets = workbook.worksheets
        max_sheets = common.MAX_SHEETS
        sections: list[str] = []
        length = -len(_SECTION_SEPARATOR)
        for sheet in worksheets[:max_sheets]:
            # The declared dimension, read before it is dropped: None without
            # one or for a rows-only one ("1:1").
            declared_columns = sheet.max_column
            sheet.reset_dimensions()
            # One column past the cap: enough for the columns note.
            read_columns = common.MAX_TABLE_COLUMNS + 1
            rows = sheet.iter_rows(values_only=True, max_col=read_columns)
            table = table_text(
                itertools.islice(rows, MAX_SHEET_ROW_INDEX),
                columns_cut=isinstance(declared_columns, int) and declared_columns > read_columns,
            )
            section = f"## {clean_cell(sheet.title)}\n\n{table or _EMPTY_SHEET}"
            length += len(_SECTION_SEPARATOR) + len(section)
            if length > common.MAX_TEXT_CHARS:
                raise ConversionError("text_too_large")
            sections.append(section)
        if len(worksheets) > max_sheets:
            sections.append(f"[Only the first {max_sheets} sheets are included.]")
        return _SECTION_SEPARATOR.join(sections)
    finally:
        workbook.close()


def convert_csv(path: Path, writer: PartWriter, options: ConversionOptions) -> int | None:
    """Write the CSV file as one Markdown table in one text part.

    The delimiter is sniffed among ``,`` ``;`` tab ``|`` from the first 8 KiB
    (``csv.excel`` when sniffing fails).

    Returns:
        None: CSV files have no pages.

    Raises:
        ConversionError: ``corrupted_file`` for bytes that aren't UTF-8 or a
            CSV error (a field over the csv module's limit);
            ``text_too_large`` for more than ``MAX_TEXT_CHARS`` characters.
    """
    with path.open(encoding="utf-8-sig", newline="") as file:
        try:
            sample = file.read(_CSV_SAMPLE_CHARS)
            # Seeking to 0 resets the decoder, so the BOM is skipped again.
            file.seek(0)
            dialect: type[csv.Dialect]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=_CSV_DELIMITERS)
            except csv.Error:
                dialect = csv.excel
            text = table_text(csv.reader(file, dialect))
        except (csv.Error, UnicodeDecodeError):
            raise ConversionError("corrupted_file") from None
    writer.add_text(text)
    return None
