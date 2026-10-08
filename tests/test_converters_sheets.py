"""Tests for ``admino.converters.sheets`` (GH-188 contract section 4.4): XLSX and CSV.

``convert_xlsx`` / ``convert_csv`` are called in-process with a real
``PartWriter``; the part file is read back as bytes. Workbooks are built with
openpyxl or, where openpyxl can't write it (cached formula values, odd sheet
names), as hand-written OOXML with ``zipfile``. Stored uploads have no
extension (``<root>/<org>/<id>``), so no fixture file has one either (openpyxl
refuses a *path* without an xlsx extension: the converter must still read it).
What these tests pin down:

- XLSX: one section per worksheet in workbook order, ``## <clean_cell(title)>``,
  a blank line, then the Markdown table (first row the header); sections joined
  by a blank line; one text part, ``page`` None, the converter returns None.
- ``format_cell``: None, bool (before int), int, integral float under 1e15
  (also negative), other floats and big floats as ``repr`` (the 1e15 boundary
  both sides, ``abs`` applied), datetime at midnight as a date, any other
  datetime (microseconds included) with a space separator, date, time,
  timedelta, str as is, anything else ``str()``; and end to end through an
  openpyxl workbook (every type read back from cells).
- Cached formula values are used; a formula without a cached value is empty.
- Table rules: cells cleaned (``|`` escaped, line breaks to spaces, the
  1000-character cut with an ellipsis), rows whose cells are all empty after
  cleaning dropped, leading and middle empty columns kept, trailing empty
  columns trimmed; ``(empty sheet)`` for a sheet without a non-empty row;
  chartsheets skipped and not counted toward the sheet cap.
- Caps with their exact notes, monkeypatched small (read at call time): rows
  (the header counts; trailing empty rows are no reason for the note), columns
  (only cells of KEPT rows count), both (rows note first, one per line after a
  blank line), sheets; plus the default 1000-row note on a real CSV.
- Sheet titles with a pipe, a line break, a tab and 1200 characters.
- Failures: not a ZIP, a ZIP without a workbook, a broken workbook part and a
  worksheet that breaks while its rows are read -> ``corrupted_file``; the
  OOXML guard refusing the archive (``MAX_OOXML_ENTRIES`` monkeypatched) ->
  ``archive_too_large``; nothing written in every case. No open file handle on
  the upload is left after a success or a failure.
- CSV: delimiter sniffed among ``,`` ``;`` tab ``|``, comma when sniffing fails,
  the BOM removed, quoted fields with delimiters and line breaks, the same table
  rules and notes, no heading; a field over the csv module's limit and bytes
  that aren't UTF-8 -> ``corrupted_file``; empty content (no bytes, a BOM
  alone, only empty rows) -> no part.
- Bounded work (contract 12.3, audit L-2/L-3; Decision 15 "tables stop at the
  text cap while being built; CSV sniffing reads 8 KiB; XLSX reads at most
  1,048,576 rows per sheet"):
  - ``table_text`` counts the characters of the cells it keeps (after
    ``clean_cell``, so an escaped pipe counts twice; cells beyond the column
    cap don't count) and raises ``text_too_large`` as soon as the total
    passes ``MAX_TEXT_CHARS`` (monkeypatched small): exactly the cap passes,
    and a row generator shows it stops at the crossing row. XLSX and CSV stop
    reading rows there too (rows pulled from openpyxl / ``csv.reader``
    counted), nothing written.
  - CSV sniffing gets the first 8 KiB of the decoded text (``csv.Sniffer
    .sniff`` spied) and still finds the delimiter; the auditor's hostile 64 KiB
    sample (``'; "a'`` repeated, quadratic in the sniffer's regexes) converts
    in under 3 seconds.
  - XLSX iterates at most ``MAX_SHEET_ROW_INDEX`` = 1_048_576 rows per sheet
    (the constant may live in ``sheets`` or ``common``; it is read at call
    time): with it monkeypatched to 1000, a row at ``r="1000"`` is kept and
    one at ``r="1001"`` is never reached (openpyxl pads missing row indexes
    with empty rows); at the real limit a row at ``r="1000000000000"`` stops
    after 1,048,576 rows in well under the timeout.
- Bounded per-row and per-cell work (contract 13.2, parser re-audit LN-1; the
  converter's ``format_cell`` / ``clean_cell`` spied in ``sheets`` and ``common``):
  - (a) A row whose values are all None or ``""`` is skipped before any per-cell
    work (``table_text`` directly: those values are never formatted or cleaned).
    The auditor's file, 67,000 rows of one value-less cell in column ZZZ (read
    full width, openpyxl makes each row 18,278 values wide), converts to
    ``(empty sheet)`` in under 10 s without a single ``format_cell`` call, and so
    does Excel's maximum of 1,048,576 such rows (a 62 KB upload; full width,
    openpyxl's padding alone took about 96 s). Revised 13.2 (a): XLSX rows are
    read ``MAX_TABLE_COLUMNS + 1`` columns wide (openpyxl's ``max_col``; the
    widths of the rows openpyxl produces are recorded), so a kept row formats and
    cleans at most that many values; a None-only tail out to column ZZZ is no
    reason for the columns note; content in column 51 fires it, while content
    further right with column 51 empty (column 60) is never read: its text is
    absent, and in a sheet without a ``<dimension>`` there is no note either
    (the remaining trade-off, see GH-281 below). Cells beyond the cap that hold no
    value at all don't fire the note (the 13.2 rule: a value other than None,
    checked unformatted, so the no-note case uses value-less cells, not
    whitespace).
  - (b) Within one ``table_text`` call a ``str`` longer than ``MAX_CELL_CHARS``
    is cleaned once per distinct value (equal values share it; a second call
    cleans it again), with output identical to cleaning every cell (pipes
    escaped after the cut, fresh values per row not mixed up). The memo is
    bounded (the addendum: an LRU of at most 1024 values): a value is cleaned
    again after 1024 other long values. The auditor's file, one
    9,000,000-character shared string in 2,050 cells, converts in under 10 s
    with that string cleaned once.
  Spies stop the slow paths early (a ``BaseException`` past a call budget, plus
  the SIGALRM deadline), so these tests fail fast on code that does the work.
- Columns note from the declared dimension (GH-281 Decision 4, contract S1-S4):
  the note also appears when the worksheet's ``<dimension>`` ends right of
  column ``MAX_TABLE_COLUMNS + 1`` (cap read at call time; with a cap of 2,
  column 3 doesn't fire it, column 4 does): data in columns A and 60 with
  column 51 empty, openpyxl-written (openpyxl always writes the dimension) or
  hand-built with ``A1:BH1``; content right of column 51 only in rows past the
  row cap (the rows note first). Rows are still read 51 columns wide. A sheet
  without a dimension, or with a rows-only one (``1:1``), keeps only the
  column-51 rule; a malformed one is still ``corrupted_file``; one claiming
  fewer rows than the sheet has cuts none; a huge one (``A1:XFD1048576``) adds
  the note and no row or cell work (rows pulled from openpyxl counted).
  When the kept rows hold nothing in their first 50 columns
  but the columns note applies, the table text is the notes alone
  (``table_text``, so XLSX, CSV and DOCX): an XLSX section is the note, not
  ``(empty sheet)``; a sheet with nothing at all stays ``(empty sheet)``.

The converter modules are imported inside fixtures, so this file collects
before they exist and every test fails on its own.
"""

from __future__ import annotations

import csv
import datetime as dt
import gc
import os
import signal
import sys
import time
import uuid
import zipfile
from contextlib import contextmanager, suppress
from dataclasses import dataclass, field
from decimal import Decimal
from typing import TYPE_CHECKING, Any

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from collections.abc import Iterator
    from pathlib import Path
    from types import ModuleType

_ELLIPSIS = chr(0x2026)
_UTF8_BOM = b"\xef\xbb\xbf"

_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT_PREFIX = "application/vnd.openxmlformats-"
_CT_WORKBOOK = _CT_PREFIX + "officedocument.spreadsheetml.sheet.main+xml"
_CT_SHEET = _CT_PREFIX + "officedocument.spreadsheetml.worksheet+xml"
_CT_RELS = _CT_PREFIX + "package.relationships+xml"
_CT_SHARED = _CT_PREFIX + "officedocument.spreadsheetml.sharedStrings+xml"


@pytest.fixture
def common() -> ModuleType:
    import admino.converters.common as module

    return module


@pytest.fixture
def sheets() -> ModuleType:
    import admino.converters.sheets as module

    return module


@pytest.fixture
def out_dir(tmp_path: Path) -> Path:
    directory = tmp_path / "out.d"
    directory.mkdir()
    return directory


# --- fixture builders ---------------------------------------------------------


def _upload_path(tmp_path: Path) -> Path:
    return tmp_path / str(uuid.uuid4())


def _workbook(
    tmp_path: Path, tabs: list[tuple[str, list[list[object]]]], *, chart_at: int | None = None
) -> Path:
    """Save a workbook of ``(title, rows)`` worksheets; a chartsheet at ``chart_at``."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for title, rows in tabs:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    if chart_at is not None:
        chart = BarChart()
        chart.add_data(Reference(book.worksheets[0], min_col=1, min_row=1, max_row=2))
        book.create_chartsheet("Chart", chart_at).add_chart(chart)
    path = _upload_path(tmp_path)
    book.save(path)
    return path


def _sheet_xml(rows_xml: str) -> str:
    return f'<worksheet xmlns="{_NS_MAIN}"><sheetData>{rows_xml}</sheetData></worksheet>'


def _sheet_xml_with_dimension(ref: str | None, rows_xml: str) -> str:
    """A worksheet declaring ``<dimension ref="..."/>`` (none when ``ref`` is None)."""
    dimension = "" if ref is None else f'<dimension ref="{ref}"/>'
    return f'<worksheet xmlns="{_NS_MAIN}">{dimension}<sheetData>{rows_xml}</sheetData></worksheet>'


def _xlsx_by_hand(
    tmp_path: Path,
    tabs: list[tuple[str, str]],
    *,
    workbook_xml: str | None = None,
    shared_strings: list[str] | None = None,
) -> Path:
    """Write an OOXML workbook by hand: ``(XML-escaped sheet name, worksheet XML)``;
    ``shared_strings`` (XML-escaped) adds a shared string table (``t="s"`` cells)."""
    count = len(tabs)
    overrides = "".join(
        f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="{_CT_SHEET}"/>'
        for i in range(1, count + 1)
    )
    shared_rel = ""
    if shared_strings is not None:
        overrides += f'<Override PartName="/xl/sharedStrings.xml" ContentType="{_CT_SHARED}"/>'
        shared_rel = (
            f'<Relationship Id="rId{count + 1}" Type="{_NS_REL}/sharedStrings" '
            'Target="sharedStrings.xml"/>'
        )
    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        f'<Default Extension="rels" ContentType="{_CT_RELS}"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/xl/workbook.xml" ContentType="{_CT_WORKBOOK}"/>'
        f"{overrides}</Types>"
    )
    root_rels = (
        f'<Relationships xmlns="{_NS_PKG_REL}"><Relationship Id="rId1" '
        f'Type="{_NS_REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    )
    entries = "".join(
        f'<sheet name="{name}" sheetId="{i}" r:id="rId{i}"/>' for i, (name, _) in enumerate(tabs, 1)
    )
    workbook = workbook_xml or (
        f'<workbook xmlns="{_NS_MAIN}" xmlns:r="{_NS_REL}"><sheets>{entries}</sheets></workbook>'
    )
    workbook_rels = (
        f'<Relationships xmlns="{_NS_PKG_REL}">'
        + "".join(
            f'<Relationship Id="rId{i}" Type="{_NS_REL}/worksheet" '
            f'Target="worksheets/sheet{i}.xml"/>'
            for i in range(1, count + 1)
        )
        + shared_rel
        + "</Relationships>"
    )
    path = _upload_path(tmp_path)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", root_rels)
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        for i, (_, sheet_xml) in enumerate(tabs, 1):
            archive.writestr(f"xl/worksheets/sheet{i}.xml", sheet_xml)
        if shared_strings is not None:
            items = "".join(f"<si><t>{text}</t></si>" for text in shared_strings)
            archive.writestr("xl/sharedStrings.xml", f'<sst xmlns="{_NS_MAIN}">{items}</sst>')
    return path


def _inline(ref: str, text: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{text}</t></is></c>'


def _csv(tmp_path: Path, data: bytes) -> Path:
    path = _upload_path(tmp_path)
    path.write_bytes(data)
    return path


def _run(
    common: ModuleType, convert: Any, path: Path, out_dir: Path, filename: str
) -> tuple[Any, Any]:
    writer = common.PartWriter(out_dir)
    options = common.ConversionOptions(filename=filename, render_dpi=150, max_pages=100)
    return convert(path, writer, options), writer


def _xlsx_text(common: ModuleType, sheets: ModuleType, path: Path, out_dir: Path) -> str:
    _run(common, sheets.convert_xlsx, path, out_dir, "book.xlsx")
    assert sorted(p.name for p in out_dir.iterdir()) == ["part-0001.txt"]
    return (out_dir / "part-0001.txt").read_bytes().decode("utf-8")


def _csv_text(common: ModuleType, sheets: ModuleType, path: Path, out_dir: Path) -> str:
    _run(common, sheets.convert_csv, path, out_dir, "data.csv")
    assert sorted(p.name for p in out_dir.iterdir()) == ["part-0001.txt"]
    return (out_dir / "part-0001.txt").read_bytes().decode("utf-8")


def _open_paths() -> set[str]:
    """Paths of every file descriptor this process has open (Linux and macOS)."""
    paths: set[str] = set()
    for name in os.listdir("/dev/fd"):
        try:
            if sys.platform == "linux":
                paths.add(os.readlink(f"/proc/self/fd/{name}"))
            else:
                import fcntl

                raw = fcntl.fcntl(int(name), fcntl.F_GETPATH, bytes(1024))
                paths.add(raw.split(b"\0", 1)[0].decode())
        except OSError:
            continue
    return paths


# --- XLSX: sections, cells, formulas ------------------------------------------


def test_sheets_xlsx_one_section_per_worksheet_in_workbook_order(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    path = _workbook(
        tmp_path,
        [
            ("Zeta", [["Region", "Q1"], ["North", 10], ["South", 12]]),
            ("Alpha", [["Text"], ["hello"]]),
        ],
    )
    result, writer = _run(common, sheets.convert_xlsx, path, out_dir, "book.xlsx")
    text = (out_dir / "part-0001.txt").read_bytes().decode("utf-8")
    assert (result, text, [(p.type, p.file, p.page) for p in writer.parts]) == (
        None,
        "## Zeta\n\n"
        "| Region | Q1 |\n| --- | --- |\n| North | 10 |\n| South | 12 |\n\n"
        "## Alpha\n\n"
        "| Text |\n| --- |\n| hello |",
        [("text", "part-0001.txt", None)],
    )


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, ""),
        (True, "TRUE"),
        (False, "FALSE"),
        (42, "42"),
        (3.0, "3"),
        (-2.0, "-2"),
        (999_999_999_999_999.0, "999999999999999"),
        (1e15, "1000000000000000.0"),
        (-1e15, "-1000000000000000.0"),
        (2.5, "2.5"),
        (0.1 + 0.2, "0.30000000000000004"),
        (1e20, "1e+20"),
        (dt.datetime(2026, 10, 7), "2026-10-07"),
        (dt.datetime(2026, 10, 7, 13, 45, 30), "2026-10-07 13:45:30"),
        (dt.datetime(2026, 10, 7, 0, 0, 0, 5), "2026-10-07 00:00:00.000005"),
        (dt.date(2026, 1, 2), "2026-01-02"),
        (dt.time(8, 30), "08:30:00"),
        (dt.timedelta(days=1, hours=2), "1 day, 2:00:00"),
        ("  raw | text ", "  raw | text "),
        (Decimal("1.50"), "1.50"),
    ],
    ids=[
        "none",
        "true",
        "false",
        "int",
        "integral-float",
        "negative-integral-float",
        "integral-float-below-1e15",
        "float-at-1e15",
        "negative-float-at-1e15",
        "fraction",
        "float-repr",
        "big-float",
        "datetime-midnight",
        "datetime-with-time",
        "datetime-with-microseconds",
        "date",
        "time",
        "timedelta",
        "str-as-is",
        "other-type-str",
    ],
)
def test_sheets_format_cell_formats_each_type(
    sheets: ModuleType, value: object, expected: str
) -> None:
    assert sheets.format_cell(value) == expected


def test_sheets_xlsx_formats_every_cell_type_read_from_the_workbook(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    header = ["bool", "int", "whole", "frac", "big", "day", "stamp", "time", "span", "none", "str"]
    values: list[object] = [
        True,
        42,
        3.0,
        2.5,
        1e20,
        dt.datetime(2026, 10, 7),
        dt.datetime(2026, 10, 7, 13, 45, 30),
        dt.time(8, 30),
        dt.timedelta(hours=26, minutes=3),
        None,
        "text",
    ]
    path = _workbook(tmp_path, [("Types", [header, values])])
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## Types\n\n"
        "| bool | int | whole | frac | big | day | stamp | time | span | none | str |\n"
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |\n"
        "| TRUE | 42 | 3 | 2.5 | 1e+20 | 2026-10-07 | 2026-10-07 13:45:30 | 08:30:00 "
        "| 1 day, 2:03:00 |  | text |"
    )


def test_sheets_xlsx_uses_cached_formula_values_and_leaves_uncached_empty(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    rows = (
        '<row r="1">'
        + _inline("A1", "Amount")
        + _inline("B1", "Double")
        + _inline("C1", "Label")
        + '</row><row r="2"><c r="A2"><f>20+1</f><v>21</v></c><c r="B2"><f>A2*2</f></c>'
        '<c r="C2" t="str"><f>"a"&amp;"b"</f><v>ab</v></c></row>'
    )
    path = _xlsx_by_hand(tmp_path, [("Calc", _sheet_xml(rows))])
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## Calc\n\n| Amount | Double | Label |\n| --- | --- | --- |\n| 21 |  | ab |"
    )


def test_sheets_xlsx_drops_empty_rows_and_trims_trailing_empty_columns(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    rows: list[list[object]] = [
        ["Name", None, "City", "   "],
        [],
        [None, None, None],
        ["  ", "\n", None, "\t"],
        ["Anna", None, "Bern", None],
        [None, "x", None, None],
    ]
    path = _workbook(tmp_path, [("People", rows)])
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## People\n\n| Name |  | City |\n| --- | --- | --- |\n| Anna |  | Bern |\n|  | x |  |"
    )


def test_sheets_xlsx_sheet_without_a_non_empty_row_is_marked_empty(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    path = _workbook(
        tmp_path,
        [("Data", [["a"]]), ("Blank", []), ("Spaces", [["  ", None], ["\t"]])],
    )
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## Data\n\n| a |\n| --- |\n\n## Blank\n\n(empty sheet)\n\n## Spaces\n\n(empty sheet)"
    )


def test_sheets_xlsx_skips_chartsheets_and_does_not_count_them(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("First", [["a"], [1]]), ("Second", [["b"]])], chart_at=1)
    monkeypatch.setattr(common, "MAX_SHEETS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## First\n\n| a |\n| --- |\n| 1 |\n\n## Second\n\n| b |\n| --- |"
    )


def test_sheets_xlsx_cleans_cells_with_pipes_line_breaks_and_long_text(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    path = _workbook(tmp_path, [("Notes", [["Note", "Long"], ["a|b\r\nc\td", "x" * 1500]])])
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## Notes\n\n| Note | Long |\n| --- | --- |\n| a\\|b c d | " + "x" * 1000 + _ELLIPSIS + " |"
    )


def test_sheets_xlsx_sheet_titles_are_cleaned_like_cells(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    tabs = [
        ("  Q1 | Plan&apos;s #1 ", _sheet_xml("")),
        ("Line&#10;Break&#9;Tab", _sheet_xml("")),
        ("T" * 1200, _sheet_xml("")),
    ]
    path = _xlsx_by_hand(tmp_path, tabs)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## Q1 \\| Plan's #1\n\n(empty sheet)\n\n"
        "## Line Break Tab\n\n(empty sheet)\n\n"
        "## " + "T" * 1000 + _ELLIPSIS + "\n\n(empty sheet)"
    )


# --- XLSX: caps and notes -----------------------------------------------------


def test_sheets_xlsx_row_cap_keeps_the_first_rows_and_adds_the_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["h"], ["r1"], [], ["r2"], ["r3"], ["r4"]])])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 3)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S\n\n| h |\n| --- |\n| r1 |\n| r2 |\n\n[Only the first 3 rows are included.]"
    )


def test_sheets_xlsx_row_cap_reached_exactly_has_no_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["h"], ["r1"], ["r2"], [], ["  "], [None, "\n"]])])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 3)
    assert _xlsx_text(common, sheets, path, out_dir) == ("## S\n\n| h |\n| --- |\n| r1 |\n| r2 |")


def test_sheets_xlsx_column_cap_drops_cells_beyond_and_adds_the_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["a", "b", "c"], ["1", "2", None]])])
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n[Only the first 2 columns are included.]"
    )


def test_sheets_xlsx_column_cap_without_content_beyond_has_no_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # Contract 13.2 (a): the note looks at the values beyond the cap unformatted (any
    # value other than None counts), so "nothing beyond the cap" is cells without a
    # value: D1 (C1 absent, padded) and C2 exist, both value-less.
    rows = (
        '<row r="1">' + _inline("A1", "a") + _inline("B1", "b") + '<c r="D1"/></row>'
        '<row r="2">' + _inline("A2", "1") + _inline("B2", "2") + '<c r="C2"/></row>'
    )
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))])
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S\n\n| a | b |\n| --- | --- |\n| 1 | 2 |"
    )


def test_sheets_xlsx_column_note_counts_kept_rows_only(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["a", "b"], ["1", "2"], ["x", "y", "z"]])])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 2)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n[Only the first 2 rows are included.]"
    )


def test_sheets_xlsx_both_caps_put_the_rows_note_first(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["a", "b", "c"], [1, 2, 3], [4, 5, 6]])])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 2)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S\n\n| a | b |\n| --- | --- |\n| 1 | 2 |\n\n"
        "[Only the first 2 rows are included.]\n"
        "[Only the first 2 columns are included.]"
    )


def test_sheets_xlsx_sheet_cap_keeps_the_first_sheets_and_adds_the_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S1", [["a"]]), ("S2", []), ("S3", [["c"]])])
    monkeypatch.setattr(common, "MAX_SHEETS", 2)
    assert _xlsx_text(common, sheets, path, out_dir) == (
        "## S1\n\n| a |\n| --- |\n\n## S2\n\n(empty sheet)\n\n"
        "[Only the first 2 sheets are included.]"
    )


# --- XLSX: failures and handles -------------------------------------------------


def _not_a_zip(tmp_path: Path) -> Path:
    path = _upload_path(tmp_path)
    path.write_bytes(b"PK\x03\x04" + b"\x00" * 60)
    return path


def _zip_without_workbook(tmp_path: Path) -> Path:
    path = _upload_path(tmp_path)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "hi")
    return path


def _broken_workbook_part(tmp_path: Path) -> Path:
    return _xlsx_by_hand(tmp_path, [("S", _sheet_xml(""))], workbook_xml="<workbook><sheets")


def _sheet_broken_while_reading(tmp_path: Path) -> Path:
    # The worksheet opens fine (its dimension comes first) and breaks in row 2.
    sheet_xml = (
        f'<worksheet xmlns="{_NS_MAIN}"><dimension ref="A1:A2"/><sheetData>'
        '<row r="1"><c r="A1"><v>1</v></c></row><row r="2"><c r="A2"'
    )
    return _xlsx_by_hand(tmp_path, [("S", sheet_xml)])


@pytest.mark.parametrize(
    "build",
    [_not_a_zip, _zip_without_workbook, _broken_workbook_part, _sheet_broken_while_reading],
    ids=["not-a-zip", "zip-without-workbook", "broken-workbook-part", "sheet-broken-in-rows"],
)
def test_sheets_xlsx_unreadable_workbook_fails_as_corrupted(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path, build: Any
) -> None:
    path = build(tmp_path)
    with pytest.raises(common.ConversionError) as caught:
        _run(common, sheets.convert_xlsx, path, out_dir, "book.xlsx")
    assert (caught.value.reason, list(out_dir.iterdir())) == ("corrupted_file", [])


def test_sheets_xlsx_archive_refused_by_the_guard_fails_as_too_large(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    path = _workbook(tmp_path, [("S", [["a"]])])
    monkeypatch.setattr(common, "MAX_OOXML_ENTRIES", 3)
    with pytest.raises(common.ConversionError) as caught:
        _run(common, sheets.convert_xlsx, path, out_dir, "book.xlsx")
    assert (caught.value.reason, list(out_dir.iterdir())) == ("archive_too_large", [])


@pytest.mark.parametrize("broken", [False, True], ids=["converted", "failed-while-reading"])
def test_sheets_xlsx_leaves_no_open_handle_on_the_upload(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path, broken: bool
) -> None:
    path = (
        _sheet_broken_while_reading(tmp_path)
        if broken
        else _workbook(tmp_path, [("S", [["a"], [1]])])
    )
    target = str(path.resolve())
    # No cyclic garbage collection while converting: an unclosed workbook would
    # otherwise be closed by the collector before the check.
    gc.disable()
    try:
        outcome = "converted"
        try:
            _run(common, sheets.convert_xlsx, path, out_dir, "book.xlsx")
        except common.ConversionError as exc:
            outcome = exc.reason
        left_open = target in _open_paths()
    finally:
        gc.enable()
    with path.open("rb"):
        seen_when_open = target in _open_paths()
    assert (outcome, left_open, seen_when_open) == (
        "corrupted_file" if broken else "converted",
        False,
        True,
    )


# --- CSV --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "delimiter", [",", ";", "\t", "|"], ids=["comma", "semicolon", "tab", "pipe"]
)
def test_sheets_csv_sniffs_the_delimiter(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path, delimiter: str
) -> None:
    rows = [["name", "city"], ["Anna", "Bern"], ["Luca", "Zug"]]
    data = "".join(delimiter.join(row) + "\n" for row in rows).encode()
    result, writer = _run(common, sheets.convert_csv, _csv(tmp_path, data), out_dir, "data.csv")
    text = (out_dir / "part-0001.txt").read_bytes().decode("utf-8")
    assert (result, text, [(p.type, p.file, p.page) for p in writer.parts]) == (
        None,
        "| name | city |\n| --- | --- |\n| Anna | Bern |\n| Luca | Zug |",
        [("text", "part-0001.txt", None)],
    )


def test_sheets_csv_falls_back_to_comma_when_sniffing_fails(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    # csv.Sniffer can't determine a delimiter for this sample.
    path = _csv(tmp_path, b"k;v,w\nm\n")
    assert _csv_text(common, sheets, path, out_dir) == "| k;v | w |\n| --- | --- |\n| m |  |"


def test_sheets_csv_removes_the_bom(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    # A BOM left in front of the first field would break its quoting.
    path = _csv(tmp_path, _UTF8_BOM + b'"Name, first",city\n"Doe, Jane",Bern\n')
    assert _csv_text(common, sheets, path, out_dir) == (
        "| Name, first | city |\n| --- | --- |\n| Doe, Jane | Bern |"
    )


def test_sheets_csv_quoted_fields_keep_delimiters_and_line_breaks(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    data = b'name,note\n"Doe, Jane","line1\r\nline2"\n"Pipe","a|b"\n'
    assert _csv_text(common, sheets, _csv(tmp_path, data), out_dir) == (
        "| name | note |\n| --- | --- |\n| Doe, Jane | line1 line2 |\n| Pipe | a\\|b |"
    )


def test_sheets_csv_applies_the_table_rules_and_both_notes(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    long_text = "y" * 1200
    data = (f"a,b,c,d\n\n , ,\t,\n1,{long_text},3,4\n5,6,7,8\n9,10,11,12\n").encode()
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 3)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 3)
    assert _csv_text(common, sheets, _csv(tmp_path, data), out_dir) == (
        "| a | b | c |\n| --- | --- | --- |\n"
        f"| 1 | {'y' * 1000}{_ELLIPSIS} | 3 |\n"
        "| 5 | 6 | 7 |\n\n"
        "[Only the first 3 rows are included.]\n"
        "[Only the first 3 columns are included.]"
    )


def test_sheets_csv_default_row_cap_is_1000_with_the_note(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    data = ("id,value\n" + "".join(f"{i},v{i}\n" for i in range(1, 1001))).encode()
    lines = _csv_text(common, sheets, _csv(tmp_path, data), out_dir).split("\n")
    # Header, separator, 999 data rows, a blank line, the note.
    assert (len(lines), lines[0], lines[-3], lines[-2:]) == (
        1003,
        "| id | value |",
        "| 999 | v999 |",
        ["", "[Only the first 1000 rows are included.]"],
    )


def test_sheets_csv_field_over_the_csv_limit_fails_as_corrupted(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    data = b"a,b\n" + b"x" * 200_000 + b",1\n"
    with pytest.raises(common.ConversionError) as caught:
        _run(common, sheets.convert_csv, _csv(tmp_path, data), out_dir, "data.csv")
    assert (caught.value.reason, list(out_dir.iterdir())) == ("corrupted_file", [])


def test_sheets_csv_non_utf8_bytes_fail_as_corrupted(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    with pytest.raises(common.ConversionError) as caught:
        _run(common, sheets.convert_csv, _csv(tmp_path, b"a,b\ncaf\xe9,1\n"), out_dir, "d.csv")
    assert (caught.value.reason, list(out_dir.iterdir())) == ("corrupted_file", [])


@pytest.mark.parametrize(
    "data",
    [b"", _UTF8_BOM, b"\n\n", b",,\n , \t\n"],
    ids=["no-bytes", "bom-only", "blank-lines", "empty-cells-only"],
)
def test_sheets_csv_without_content_writes_no_part(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path, data: bytes
) -> None:
    result, writer = _run(common, sheets.convert_csv, _csv(tmp_path, data), out_dir, "e.csv")
    assert (result, list(out_dir.iterdir()), writer.parts) == (None, [], [])


# --- bounded work: the text cap while building, the CSV sample, the XLSX row index -------


class _RanPastLimitError(BaseException):
    """Raised when more rows are pulled than allowed (not an Exception, so the
    converter's error mapping can't swallow it)."""


class _TooSlowError(BaseException):
    """Raised by the deadline below."""


@contextmanager
def _deadline(seconds: float) -> Iterator[None]:
    """Interrupt the code under test after ``seconds`` (SIGALRM, main thread)."""

    def expired(signum: int, frame: object) -> None:
        raise _TooSlowError(f"still running after {seconds} s")

    previous = signal.signal(signal.SIGALRM, expired)
    signal.setitimer(signal.ITIMER_REAL, seconds)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


def _count_xlsx_rows(
    monkeypatch: pytest.MonkeyPatch, *, stop_after: int | None = None
) -> list[int]:
    """Count the rows openpyxl's read-only worksheets produce (padding rows included)."""
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    real = ReadOnlyWorksheet._cells_by_row
    pulled = [0]

    def counted(self: Any, *args: Any, **kwargs: Any) -> Iterator[Any]:
        for row in real(self, *args, **kwargs):
            pulled[0] += 1
            if stop_after is not None and pulled[0] > stop_after:
                raise _RanPastLimitError(pulled[0])
            yield row

    monkeypatch.setattr(ReadOnlyWorksheet, "_cells_by_row", counted)
    return pulled


def _count_csv_rows(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """Count the rows ``csv.reader`` produces."""
    real = csv.reader
    pulled = [0]

    def counted(*args: Any, **kwargs: Any) -> Iterator[list[str]]:
        for row in real(*args, **kwargs):
            pulled[0] += 1
            yield row

    monkeypatch.setattr(csv, "reader", counted)
    return pulled


def test_sheets_table_text_raises_text_too_large_as_soon_as_kept_cells_pass_the_cap(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(common, "MAX_TEXT_CHARS", 100)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    # Exactly 100 kept characters: the "y" cell is beyond the column cap and doesn't count.
    at_cap = sheets.table_text([["x" * 25, "x" * 25], ["x" * 25, "x" * 25, "y" * 30]])
    pulled = 0

    def rows() -> Iterator[list[str]]:
        nonlocal pulled
        for _ in range(1000):
            pulled += 1
            yield ["|" * 15]  # 30 kept characters once every pipe is escaped

    with pytest.raises(common.ConversionError) as caught:
        sheets.table_text(rows())

    x25 = "x" * 25
    assert (at_cap, caught.value.reason, pulled) == (
        f"| {x25} | {x25} |\n| --- | --- |\n| {x25} | {x25} |\n\n"
        "[Only the first 2 columns are included.]",
        "text_too_large",
        4,
    )


@pytest.mark.parametrize("kind", ["xlsx", "csv"])
def test_sheets_conversion_stops_reading_rows_once_the_text_cap_is_passed(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
    kind: str,
) -> None:
    rows: list[list[object]] = [["x" * 100] for _ in range(1000)]
    if kind == "xlsx":
        path = _workbook(tmp_path, [("S", rows)])
        pulled = _count_xlsx_rows(monkeypatch)
        convert = sheets.convert_xlsx
    else:
        path = _csv(tmp_path, "".join(f"{row[0]}\n" for row in rows).encode())
        pulled = _count_csv_rows(monkeypatch)
        convert = sheets.convert_csv
    monkeypatch.setattr(common, "MAX_TEXT_CHARS", 1000)  # passed by the 11th row

    with pytest.raises(common.ConversionError) as caught:
        _run(common, convert, path, out_dir, f"data.{kind}")

    assert (caught.value.reason, pulled[0] <= 12, list(out_dir.iterdir())) == (
        "text_too_large",
        True,
        [],
    )


def test_sheets_csv_sniffs_only_the_first_8_kib(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    text = "name;city\n" + "".join(f"person{i};town{i}\n" for i in range(2000))
    assert len(text) > 4 * 8192
    samples: list[str] = []
    real = csv.Sniffer.sniff

    def sniff(self: csv.Sniffer, sample: str, delimiters: str | None = None) -> Any:
        samples.append(sample)
        return real(self, sample, delimiters)

    monkeypatch.setattr(csv.Sniffer, "sniff", sniff)

    first_line = _csv_text(common, sheets, _csv(tmp_path, text.encode()), out_dir).split("\n")[0]

    assert (
        [len(sample) for sample in samples],
        all(text.startswith(sample) for sample in samples),
        first_line,
    ) == ([8192], True, "| name | city |")


def test_sheets_csv_hostile_sniffer_sample_converts_quickly(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    # The auditor's L-3 probe: csv.Sniffer's regexes are quadratic on this
    # pattern (about 15 s on a 64 KiB sample, about 0.25 s on 8 KiB).
    pattern = '; "a'
    path = _csv(tmp_path, (pattern * (64 * 1024 // len(pattern))).encode())

    start = time.perf_counter()
    with suppress(common.ConversionError):
        _run(common, sheets.convert_csv, path, out_dir, "hostile.csv")

    assert time.perf_counter() - start < 3.0


def _row_limit_modules(common: ModuleType, sheets: ModuleType) -> list[ModuleType]:
    """Where ``MAX_SHEET_ROW_INDEX`` is defined (sheets or common)."""
    return [module for module in (sheets, common) if hasattr(module, "MAX_SHEET_ROW_INDEX")]


def _far_row_workbook(tmp_path: Path, index: int) -> Path:
    """A sheet with ``head`` in row 1 and ``tail`` in row ``index`` (nothing between)."""
    rows = (
        '<row r="1">'
        + _inline("A1", "head")
        + f'</row><row r="{index}">'
        + _inline(f"A{index}", "tail")
        + "</row>"
    )
    return _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))])


def test_sheets_xlsx_reads_at_most_max_sheet_row_index_rows(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    modules = _row_limit_modules(common, sheets)
    values = {module.MAX_SHEET_ROW_INDEX for module in modules}
    for module in modules:
        monkeypatch.setattr(module, "MAX_SHEET_ROW_INDEX", 1000)
    texts = []
    for index in (1000, 1001):
        out = tmp_path / f"out-{index}"
        out.mkdir()
        texts.append(_xlsx_text(common, sheets, _far_row_workbook(tmp_path, index), out))

    assert (values, texts) == (
        {1_048_576},
        ["## S\n\n| head |\n| --- |\n| tail |", "## S\n\n| head |\n| --- |"],
    )


def test_sheets_xlsx_huge_row_index_stops_after_the_last_excel_row(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # The auditor's L-3 probe: openpyxl yields one empty row per missing index,
    # so r="1000000000000" would otherwise run until the conversion timeout.
    path = _far_row_workbook(tmp_path, 10**12)
    pulled = _count_xlsx_rows(monkeypatch, stop_after=1_048_576)

    start = time.perf_counter()
    with _deadline(60):
        text = _xlsx_text(common, sheets, path, out_dir)
    seconds = time.perf_counter() - start

    assert (text, pulled[0], seconds < 20) == ("## S\n\n| head |\n| --- |", 1_048_576, True)


# --- bounded per-row and per-cell work (contract 13.2, parser re-audit LN-1) ----------


@dataclass
class _CellCalls:
    """What ``format_cell`` and ``clean_cell`` received while converting;
    ``long_cleanings`` counts the cleanings of a value over 1,000,000 characters."""

    formatted: list[object] = field(default_factory=list)
    cleaned: list[str] = field(default_factory=list)
    long_cleanings: int = 0


def _spy_cells(
    monkeypatch: pytest.MonkeyPatch,
    common: ModuleType,
    sheets: ModuleType,
    *,
    format_budget: int | None = None,
    long_clean_budget: int | None = None,
) -> _CellCalls:
    """Record every ``format_cell`` / ``clean_cell`` call (``sheets.clean_cell`` and
    ``common.clean_cell`` alike), each still doing its work. Past ``format_budget``
    formatted values, or ``long_clean_budget`` cleanings of a value over 1,000,000
    characters, the spy raises (a BaseException, so the converter's error mapping
    can't swallow it): code that does the unbounded work fails at once instead of
    running for minutes."""
    real_format = sheets.format_cell
    real_clean = common.clean_cell
    calls = _CellCalls()

    def format_spy(value: object) -> str:
        calls.formatted.append(value)
        if format_budget is not None and len(calls.formatted) > format_budget:
            raise _RanPastLimitError(f"format_cell called more than {format_budget} times")
        return str(real_format(value))

    def clean_spy(value: str) -> str:
        calls.cleaned.append(value)
        if len(value) > 1_000_000:
            calls.long_cleanings += 1
        if long_clean_budget is not None and calls.long_cleanings > long_clean_budget:
            raise _RanPastLimitError(f"a huge value cleaned {calls.long_cleanings} times")
        return str(real_clean(value))

    monkeypatch.setattr(sheets, "format_cell", format_spy)
    monkeypatch.setattr(sheets, "clean_cell", clean_spy)
    monkeypatch.setattr(common, "clean_cell", clean_spy)
    return calls


def test_sheets_table_text_rows_of_only_none_or_empty_strings_are_never_formatted(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    # 13.2 (a): a row whose values are all None or "" is skipped before any per-cell work.
    calls = _spy_cells(monkeypatch, common, sheets)

    text = sheets.table_text([["h"], [None, None, None], ["", "", None, ""], [], ["x"]])

    assert (text, set(calls.formatted) <= {"h", "x"}, set(calls.cleaned) <= {"h", "x"}) == (
        "| h |\n| --- |\n| x |",
        True,
        True,
    )


def test_sheets_xlsx_auditor_wide_empty_rows_convert_fast_without_cell_work(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # The auditor's LN-1 (a) file (about 5 KB): 67,000 rows of one value-less cell in
    # column ZZZ. openpyxl makes every row 18,278 values wide; formatting and cleaning
    # each of them took about 120 s, the whole conversion timeout.
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml('<row><c r="ZZZ1"/></row>' * 67_000))])
    calls = _spy_cells(monkeypatch, common, sheets, format_budget=0)

    start = time.perf_counter()
    with _deadline(60):
        text = _xlsx_text(common, sheets, path, out_dir)
    seconds = time.perf_counter() - start

    assert (
        text,
        calls.formatted,
        [value for value in calls.cleaned if value != "S"],
        seconds < 10,
    ) == ("## S\n\n(empty sheet)", [], [], True)


def _openpyxl_seconds(path: Path, columns: int) -> float:
    """How long openpyxl alone takes to produce the rows of the first sheet of ``path``,
    ``columns`` wide: the XML parsing no converter can skip."""
    with path.open("rb") as file:
        book = openpyxl.load_workbook(file, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            start = time.perf_counter()
            for _row in sheet.iter_rows(values_only=True, max_col=columns):
                pass
            return time.perf_counter() - start
        finally:
            book.close()


def test_sheets_xlsx_a_million_wide_empty_rows_convert_fast(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # Revised 13.2 (a) budget: the row above, Excel's maximum of 1,048,576 times (a 62 KB
    # upload). Read full width, openpyxl's padding alone takes about 96 s. Read 51 columns
    # wide, openpyxl's XML parsing is the floor: about 3 s on an M2, 7 s under coverage's
    # tracer (`make check`), where a fixed 10 s can't hold. So the conversion may take at
    # most three times that floor (openpyxl timed on an eighth of the rows, times 8) plus
    # 1 s; about 5 s here (9.5 s traced). The spy (no format_cell call at all) and the
    # deadline end slow code early.
    row = '<row><c r="ZZZ1"/></row>'
    eighth = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(row * (1_048_576 // 8)))])
    floor = 8 * _openpyxl_seconds(eighth, common.MAX_TABLE_COLUMNS + 1)
    budget = 3 * floor + 1
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(row * 1_048_576))])
    calls = _spy_cells(monkeypatch, common, sheets, format_budget=0)

    start = time.perf_counter()
    with _deadline(budget + 5):
        text = _xlsx_text(common, sheets, path, out_dir)
    seconds = time.perf_counter() - start

    assert (path.stat().st_size < 100_000, text, calls.formatted, seconds < budget) == (
        True,
        "## S\n\n(empty sheet)",
        [],
        True,
    ), f"{seconds:.1f} s for a budget of {budget:.1f} s"


def test_sheets_xlsx_wide_row_formats_at_most_the_column_cap_plus_one_values(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # Two kept rows reach out to column ZZZ with a value-less cell (wide empty rows
    # between them): each formats and cleans at most MAX_TABLE_COLUMNS + 1 values, and
    # a None-only tail is no reason for the columns note.
    rows = (
        '<row r="1">'
        + _inline("A1", "a")
        + '<c r="ZZZ1"/></row>'
        + "".join(f'<row r="{r}"><c r="ZZZ{r}"/></row>' for r in range(2, 6))
        + '<row r="6">'
        + _inline("A6", "b")
        + '<c r="ZZZ6"/></row>'
    )
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))])
    per_row = common.MAX_TABLE_COLUMNS + 1
    calls = _spy_cells(monkeypatch, common, sheets, format_budget=2 * per_row)

    with _deadline(60):
        text = _xlsx_text(common, sheets, path, out_dir)

    cell_cleanings = [value for value in calls.cleaned if value != "S"]
    assert (text, len(calls.formatted) <= 2 * per_row, len(cell_cleanings) <= 2 * per_row) == (
        "## S\n\n| a |\n| --- |\n| b |",
        True,
        True,
    )


def _xlsx_row_widths(monkeypatch: pytest.MonkeyPatch) -> list[int]:
    """Record the width of every row openpyxl's read-only worksheets produce, i.e. what
    the converter receives (openpyxl pads a row to ``max_col``, else to its last cell)."""
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    real = ReadOnlyWorksheet._cells_by_row
    widths: list[int] = []

    def measured(self: Any, *args: Any, **kwargs: Any) -> Iterator[Any]:
        for row in real(self, *args, **kwargs):
            widths.append(len(row))
            yield row

    monkeypatch.setattr(ReadOnlyWorksheet, "_cells_by_row", measured)
    return widths


def _read_at_most(widths: list[int], columns: int) -> bool:
    """Rows were read, none of them wider than ``columns``."""
    return bool(widths) and max(widths) <= columns


def test_sheets_xlsx_content_beyond_column_51_with_column_51_empty_is_not_read(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # Revised 13.2 (a): rows are read MAX_TABLE_COLUMNS + 1 = 51 columns wide, so with
    # column 51 empty the content in column 60 is never seen: its text is absent and at
    # most 51 values are formatted. This hand-built sheet declares no <dimension>, so only
    # the column-51 rule applies and there is no columns note (GH-281 Decision 4: the
    # remaining trade-off; a declared dimension right of column 51 fires it, see below).
    assert (get_column_letter(51), get_column_letter(60)) == ("AY", "BH")
    rows = '<row r="1">' + _inline("A1", "a") + _inline("BH1", "beyond") + "</row>"
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))])
    per_row = common.MAX_TABLE_COLUMNS + 1
    widths = _xlsx_row_widths(monkeypatch)
    calls = _spy_cells(monkeypatch, common, sheets, format_budget=per_row)

    text = _xlsx_text(common, sheets, path, out_dir)

    assert (text, _read_at_most(widths, per_row), len(calls.formatted) <= per_row) == (
        "## S\n\n| a |\n| --- |",
        True,
        True,
    )


def test_sheets_xlsx_content_in_column_51_fires_the_columns_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # The other side of the trade-off: content in column 51 (AY, the one column read
    # beyond the cap of 50) fires the note, though column 60 (BH) is still not read.
    rows = '<row r="1">' + _inline("A1", "a") + _inline("AY1", "y") + _inline("BH1", "z") + "</row>"
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))])
    widths = _xlsx_row_widths(monkeypatch)

    text = _xlsx_text(common, sheets, path, out_dir)

    assert (text, _read_at_most(widths, common.MAX_TABLE_COLUMNS + 1)) == (
        "## S\n\n| a |\n| --- |\n\n[Only the first 50 columns are included.]",
        True,
    )


# --- XLSX columns note from the declared dimension (GH-281 Decision 4) ------------------

_COLUMNS_NOTE = "[Only the first 50 columns are included.]"


@pytest.mark.parametrize("build", ["openpyxl-written", "hand-built-dimension"])
def test_sheets_xlsx_content_in_column_60_with_column_51_empty_adds_the_columns_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
    build: str,
) -> None:
    # Data in column A and in column 60 (BH), column 51 (AY) empty: the sheet's declared
    # dimension (A1:BH1, which openpyxl writes itself) ends right of column 51, so the
    # note appears, while rows are still read 51 columns wide (BH's text never read).
    if build == "openpyxl-written":
        path = _workbook(tmp_path, [("S", [["a", *[None] * 58, "beyond"]])])
    else:
        rows = '<row r="1">' + _inline("A1", "a") + _inline("BH1", "beyond") + "</row>"
        path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml_with_dimension("A1:BH1", rows))])
    per_row = common.MAX_TABLE_COLUMNS + 1
    widths = _xlsx_row_widths(monkeypatch)
    calls = _spy_cells(monkeypatch, common, sheets, format_budget=per_row)

    text = _xlsx_text(common, sheets, path, out_dir)

    assert (text, _read_at_most(widths, per_row), len(calls.formatted) <= per_row) == (
        f"## S\n\n| a |\n| --- |\n\n{_COLUMNS_NOTE}",
        True,
        True,
    )


def test_sheets_xlsx_dimension_covers_rows_past_the_row_cap_and_its_note_follows_the_rows_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # Column 60 holds content only in row 4, past the row cap of 2 (never read): the
    # dimension (A1:BH4) still fires the columns note, after the rows note.
    path = _workbook(tmp_path, [("S", [["h"], ["r1"], ["r2"], [*[None] * 59, "far"]])])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 2)

    assert _xlsx_text(common, sheets, path, out_dir) == (
        f"## S\n\n| h |\n| --- |\n| r1 |\n\n[Only the first 2 rows are included.]\n{_COLUMNS_NOTE}"
    )


def test_sheets_xlsx_dimension_rule_reads_the_column_cap_at_call_time(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    # Cap 2: a dimension ending at column 3 (cap + 1, the column rule (a) reads) is no
    # reason for the note, column 4 is. A dimension without a column ("1:1") or none at
    # all leaves rule (a) alone; a malformed one still fails at load. Every dimension
    # claims row 1 only while the sheet has two rows: it never cuts rows off (it is
    # dropped before the rows are read).
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    rows = '<row r="1">' + _inline("A1", "a") + '</row><row r="2">' + _inline("A2", "b") + "</row>"
    refs = {
        "ends-at-cap-plus-1": "A1:C1",
        "ends-at-cap-plus-2": "A1:D1",
        "rows-only": "1:1",
        "no-dimension": None,
        "malformed": "A1:ZZZZ1",
    }
    outcomes = {}
    for label, ref in refs.items():
        out = tmp_path / f"out-{label}"
        out.mkdir()
        path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml_with_dimension(ref, rows))])
        try:
            outcomes[label] = _xlsx_text(common, sheets, path, out)
        except common.ConversionError as exc:
            outcomes[label] = exc.reason

    table = "## S\n\n| a |\n| --- |\n| b |"
    assert outcomes == {
        "ends-at-cap-plus-1": table,
        "ends-at-cap-plus-2": f"{table}\n\n[Only the first 2 columns are included.]",
        "rows-only": table,
        "no-dimension": table,
        "malformed": "corrupted_file",
    }


@pytest.mark.parametrize("build", ["openpyxl-written", "hand-built-dimension"])
def test_sheets_xlsx_sheet_with_content_only_right_of_column_51_is_the_columns_note_alone(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path, build: str
) -> None:
    # Nothing in the 51 columns read, the dimension ends at column 60: the section is the
    # note alone, not "(empty sheet)" and not an empty table. A sheet with nothing at all
    # (and no wide dimension) stays "(empty sheet)".
    if build == "openpyxl-written":
        path = _workbook(tmp_path, [("S", [[*[None] * 59, "far"]]), ("Blank", [])])
    else:
        far = '<row r="1">' + _inline("BH1", "far") + "</row>"
        path = _xlsx_by_hand(
            tmp_path,
            [
                ("S", _sheet_xml_with_dimension("A1:BH1", far)),
                ("Blank", _sheet_xml_with_dimension("A1", "")),
            ],
        )

    assert _xlsx_text(common, sheets, path, out_dir) == (
        f"## S\n\n{_COLUMNS_NOTE}\n\n## Blank\n\n(empty sheet)"
    )


def test_sheets_xlsx_huge_declared_dimension_only_adds_the_columns_note(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # A1:XFD1048576 with two cells: the note, and no extra row or cell work (openpyxl
    # would pad up to the declared last row if the dimension were kept for reading).
    rows = '<row r="1">' + _inline("A1", "a") + '</row><row r="2">' + _inline("B2", "b") + "</row>"
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml_with_dimension("A1:XFD1048576", rows))])
    pulled = _count_xlsx_rows(monkeypatch, stop_after=10)
    widths = _xlsx_row_widths(monkeypatch)

    with _deadline(30):
        text = _xlsx_text(common, sheets, path, out_dir)

    assert (text, pulled[0], _read_at_most(widths, common.MAX_TABLE_COLUMNS + 1)) == (
        f"## S\n\n| a |  |\n| --- | --- |\n|  | b |\n\n{_COLUMNS_NOTE}",
        2,
        True,
    )


def test_sheets_table_text_content_only_beyond_the_column_cap_gives_the_notes_alone(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Decision 4: no kept cell holds content in the first 50 columns (a whitespace cell
    # cleans to nothing) but the columns note applies: the notes alone, no empty table.
    beyond_only = sheets.table_text([[None] * 50 + ["x"]])
    blank_within = sheets.table_text([["  ", *[None] * 49, "x"]])
    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 1)
    with_rows_note = sheets.table_text([[None] * 50 + ["x"], [None] * 50 + ["y"]])

    assert (beyond_only, blank_within, with_rows_note) == (
        _COLUMNS_NOTE,
        _COLUMNS_NOTE,
        f"[Only the first 1 rows are included.]\n{_COLUMNS_NOTE}",
    )


def test_sheets_csv_content_only_beyond_the_column_cap_is_the_columns_note_alone(
    common: ModuleType, sheets: ModuleType, tmp_path: Path, out_dir: Path
) -> None:
    data = ("," * 50 + "x\n" + "," * 50 + "y\n").encode()

    assert _csv_text(common, sheets, _csv(tmp_path, data), out_dir) == _COLUMNS_NOTE


def test_sheets_table_text_cleans_each_distinct_long_value_once_per_call(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    # 13.2 (b), small: values over MAX_CELL_CHARS (5 here) are cleaned once per
    # distinct value within a call ("twin" is equal to "shared", another object), again
    # in the next call, and the output is what cleaning every cell gives: cut, then the
    # pipe escaped; a fresh long value per row keeps its own text.
    monkeypatch.setattr(common, "MAX_CELL_CHARS", 5)
    shared = "".join(["ab|c", "defgh"])
    twin = "".join(["ab|cd", "efgh"])
    assert (shared == twin, shared is twin) == (True, False)
    calls = _spy_cells(monkeypatch, common, sheets)

    def rows() -> Iterator[list[str]]:
        yield ["head", "tail"]
        yield [shared, twin]
        for i in range(3):
            yield [f"row{i}-" + "p" * 8, shared]

    first = sheets.table_text(rows())
    second = sheets.table_text(rows())

    cut = "ab\\|cd" + _ELLIPSIS
    expected = (
        "| head | tail |\n| --- | --- |\n"
        f"| {cut} | {cut} |\n"
        f"| row0-{_ELLIPSIS} | {cut} |\n"
        f"| row1-{_ELLIPSIS} | {cut} |\n"
        f"| row2-{_ELLIPSIS} | {cut} |"
    )
    assert (first, second, [value for value in calls.cleaned if value == shared]) == (
        expected,
        expected,
        [shared, shared],
    )


def test_sheets_table_text_long_value_memo_holds_at_most_1024_values(
    common: ModuleType, sheets: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    # 13.2 (b) addendum: the memo is an LRU of at most 1024 distinct values. "v0000-x"
    # (over MAX_CELL_CHARS = 5) is cleaned for its first cell and reused for the second;
    # after 1024 other long values it has been dropped and is cleaned again: a memo of
    # every distinct value would grow with the file.
    monkeypatch.setattr(common, "MAX_CELL_CHARS", 5)
    values = [f"v{i:04d}-x" for i in range(1025)]
    first = values[0]
    calls = _spy_cells(monkeypatch, common, sheets)

    def rows() -> Iterator[list[str]]:
        yield [first, first]
        for begin in range(1, 1025, 50):
            yield values[begin : begin + 50]
        yield [first]

    text = sheets.table_text(rows())

    assert (
        [value for value in calls.cleaned if value == first],
        text.count("v0000" + _ELLIPSIS),
        text.count(_ELLIPSIS),
    ) == ([first, first], 3, 1027)


def test_sheets_xlsx_auditor_huge_shared_string_is_cleaned_once_and_converts_fast(
    common: ModuleType,
    sheets: ModuleType,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    out_dir: Path,
) -> None:
    # The auditor's LN-1 (b) file (about 15 KB): one 9,000,000-character shared string
    # referenced by 2,050 cells (41 rows of 50). Cleaning it per cell cost about 11 ms
    # each (about 22 s here, about 110 s up to the text cap).
    rows = "".join(
        f'<row r="{r}">'
        + "".join(f'<c r="{get_column_letter(c)}{r}" t="s"><v>0</v></c>' for c in range(1, 51))
        + "</row>"
        for r in range(1, 42)
    )
    path = _xlsx_by_hand(tmp_path, [("S", _sheet_xml(rows))], shared_strings=["x" * 9_000_000])
    calls = _spy_cells(monkeypatch, common, sheets, long_clean_budget=1)

    start = time.perf_counter()
    with _deadline(60):
        text = _xlsx_text(common, sheets, path, out_dir)
    seconds = time.perf_counter() - start

    line = "| " + " | ".join(["x" * 1000 + _ELLIPSIS] * 50) + " |"
    expected = "## S\n\n" + "\n".join([line, "| " + " | ".join(["---"] * 50) + " |", *[line] * 40])
    # Compared as a flag: a failing diff of two 2 MB texts would take minutes.
    assert (text == expected, len(text), calls.long_cleanings, seconds < 10) == (
        True,
        len(expected),
        1,
        True,
    )
