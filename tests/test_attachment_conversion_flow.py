"""End-to-end conversion of stored attachments (GH-188 contract §4 to §7).

Issue #188: "Token estimate & storage" (every processed file stores its estimated
token count; derived artifacts next to the original), "Safety limits" (parsing in
a worker process, per-file timeout, failure reason codes), "From #187 — Page
limit" (a PDF over the platform's ``max_pages_per_file`` ends
``failed(too_many_pages)``), "Tests" (programmatic fixtures per type, scanned and
mixed PDFs, encrypted files, timeout, page markers), Decisions 2, 3, 4, 11, 12.

Each test stores an ``uploaded`` row in FakeDb plus its file under an attachments
root in ``tmp_path`` and runs ``process_attachment`` (or the ``ProcessingPool``)
with the DEFAULT processor, so the real ``convert_stored_file`` starts the real
worker process (``python -m admino.converters.worker``). The worker's PYTHONPATH
is set to the ``src`` directory this test process imports admino from, so the
child runs the same tree.

What these tests pin down:

- A 3-page mixed PDF (text, scanned, text): ``ready``, ``page_count`` 3, a
  ``token_estimate`` > 0 equal to the manifest's (and to the sum of its parts);
  ``<id>.d`` holds exactly ``manifest.json``, ``part-0001.txt``,
  ``part-0002.jpg``, ``part-0003.txt``; the text part starts with the page marker
  built from the row's file name (``[<file> — page 1]``), the image part is
  labelled ``[<file> — page 2]``.
- A PDF with more pages than the stored platform ``max_pages_per_file``:
  ``failed(too_many_pages)``, no ``<id>.d`` left, the original kept.
- A DOCX, an XLSX, a CSV, a TXT and a PNG: each ``ready`` without a page count,
  with a ``token_estimate`` > 0 equal to the manifest's; one part each (text, or
  the PNG as ``image/png``); the TXT's and the PNG's estimates exactly as the
  contract's rule gives them.
- An encrypted PDF (``/Encrypt``): ``failed(password_protected)`` and no worker
  process ever starts (``runner.WORKER_ARGV`` is replaced by one that leaves a
  marker file), no ``<id>.d``.
- A worker that outlives ``runner.CONVERSION_TIMEOUT_S``: killed (its pid is gone
  afterwards, well before its own sleep ends), ``failed(conversion_timeout)``, no
  ``<id>.d``, the original kept.
- A worker that floods its stdout (1 MiB, then sleeps with stdout open; GH-281
  Decision 1): killed as soon as more than ``runner.MAX_RESULT_BYTES`` have arrived,
  ``failed(processing_error)`` long before the timeout, its pid gone, no ``<id>.d``, the
  original kept, the flood in no log record.
- ``ProcessingPool().submit(...)`` + ``join()``: the same ``ready`` outcome with
  its estimate and artifacts.
- No log record during the flows (a ready TXT, an encrypted PDF, a crashing worker
  that echoes its job, file name and path included, to stdout and stderr) holds
  the file name, a path under ``tmp_path`` or the file's text.
- A scanned page is rendered at the stored platform ``render_dpi``: 72 and 144 dpi
  give images within 1 px of the page size at those resolutions, and the JPEG on
  disk has the manifest's size.
- Derived bytes and the quota (contract §12.4, process M-3, Decision 15): a real
  conversion stores ``derived_bytes`` equal to the bytes on disk in ``<id>.d``; an
  org whose quota holds the original but not its derived files ends
  ``failed(storage_quota_exceeded)`` with no ``<id>.d`` and no estimate; derived
  files over ``common.MAX_DERIVED_BYTES`` (lowered in this server process, measured
  by the runner) end ``failed(output_too_large)`` with no ``<id>.d``.

The org gets a 1 GiB storage quota (FakeDb's orgs start at 0, and derived files now
count toward it), except where a test sets its own.

Fixtures are built in memory (hand-written PDF bytes, python-docx, openpyxl,
Pillow); no binary in the repo. The new modules are imported lazily, so the file
collects before GH-188 and every test fails on its own.
"""

from __future__ import annotations

import asyncio
import io
import json
import logging
import math
import os
import signal
import sys
import time
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any, Final

import pytest

from admino import organizations, scoped_settings
from tests.db_fakes import ORG_ID, FakeDb
from tests.log_capture import configured_logging

if TYPE_CHECKING:
    import uuid
    from collections.abc import Callable
    from types import ModuleType

_JOIN_S: Final = 120.0
_EM_DASH: Final = "—"

_TXT_TEXT: Final = "Meeting notes 2026\nSend the report to the board before Friday.\n"
_TXT: Final = _TXT_TEXT.encode("utf-8")

# A minimal PDF whose trailer names an /Encrypt dictionary (#187's detection refuses
# it as password_protected).
_ENCRYPTED_PDF: Final = (
    b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    b"trailer\n<< /Root 1 0 R /Encrypt 5 0 R >>\n%%EOF\n"
)

# A worker stand-in that echoes its job (path, out_dir and file name included) to
# stdout and stderr, then crashes.
_ECHO_AND_CRASH: Final = (
    sys.executable,
    "-c",
    "import sys; job = sys.stdin.read(); sys.stdout.write(job); sys.stderr.write(job); sys.exit(3)",
)


# --- Fixture files ------------------------------------------------------------------


def _pdf(pages: list[tuple[str, float, float]]) -> bytes:
    """Hand-written PDF bytes: one page per (content, width_pt, height_pt).

    A content of ``"scan"`` is an image-only page (a 16x16 grey image XObject
    stretched over the page, no text layer); any other content is one line of
    Helvetica text.
    """
    pixels = bytes((x * 16 + y * 8) % 256 for y in range(16) for x in range(16))
    image = zlib.compress(pixels)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # the page tree, filled in below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Type /XObject /Subtype /Image /Width 16 /Height 16 /ColorSpace /DeviceGray "
        b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n"
        % len(image)
        + image
        + b"\nendstream",
    ]
    kids: list[str] = []
    for content, width, height in pages:
        page_number = len(objects) + 1
        kids.append(f"{page_number} 0 R")
        if content == "scan":
            stream = f"q {width:.0f} 0 0 {height:.0f} 0 0 cm /Im1 Do Q".encode()
        else:
            stream = f"BT /F1 12 Tf 36 {height - 72:.0f} Td ({content}) Tj ET".encode("latin-1")
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {width:.0f} {height:.0f}] "
            f"/Resources << /Font << /F1 3 0 R >> /XObject << /Im1 4 0 R >> >> "
            f"/Contents {page_number + 1} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def _mixed_pdf() -> bytes:
    """Three pages: text, scanned (image only, 2 x 3 inches), text."""
    return _pdf(
        [
            ("Quarterly revenue grew by 12 percent in region North.", 612, 792),
            ("scan", 144, 216),
            ("Second text page with more than twenty characters.", 612, 792),
        ]
    )


def _scanned_pdf() -> bytes:
    """One image-only page of 144 x 216 pt (2 x 3 inches)."""
    return _pdf([("scan", 144, 216)])


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Budget 2026", level=1)
    document.add_paragraph("The board approved the plan for the new office.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Q3"
    sheet.append(["name", "amount"])
    sheet.append(["alpha", 1])
    sheet.append(["beta", 2.5])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def _text_tokens(text: str) -> int:
    """Contract §2: one token per ASCII digit, plus one per 4 other UTF-8 bytes."""
    digits = sum(1 for char in text if "0" <= char <= "9")
    return digits + math.ceil((len(text.encode("utf-8")) - digits) / 4)


# --- The world ----------------------------------------------------------------------


@dataclass(frozen=True)
class _Flow:
    db: FakeDb
    root: Path
    chat_id: uuid.UUID


@pytest.fixture()
def ap() -> ModuleType:
    """The module under test, imported lazily."""
    import admino.attachment_processing as module

    return module


@pytest.fixture()
def flow(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> _Flow:
    """FakeDb with one org member and chat; the attachments root in tmp_path; the
    worker child imports the same admino tree as this process."""
    import admino

    root = tmp_path / "attachments"
    root.mkdir()
    monkeypatch.setattr(organizations, "ATTACHMENTS_ROOT", root)
    monkeypatch.setenv("PYTHONPATH", str(Path(admino.__file__).resolve().parents[1]))
    db = FakeDb()
    user_id = db.add_account(org_id=ORG_ID)
    # Room for the originals and their derived files (GH-188 §12.4).
    db.add_org(ORG_ID, storage_quota_bytes=1024**3)
    return _Flow(db=db, root=root, chat_id=db.add_chat(user_id))


def _store(flow: _Flow, data: bytes, *, kind: str, filename: str) -> uuid.UUID:
    """An uploaded row and its stored file <root>/<org>/<id>."""
    attachment_id = flow.db.add_attachment(
        flow.chat_id, filename=filename, kind=kind, size_bytes=len(data)
    )
    path = _original(flow, attachment_id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return attachment_id


def _original(flow: _Flow, attachment_id: uuid.UUID) -> Path:
    return flow.root / str(ORG_ID) / str(attachment_id)


def _derived(flow: _Flow, attachment_id: uuid.UUID) -> Path:
    return flow.root / str(ORG_ID) / f"{attachment_id}.d"


def _gone(path: Path) -> bool:
    return not path.exists() and not path.is_symlink()


def _row(flow: _Flow, attachment_id: uuid.UUID) -> tuple[Any, Any, Any, Any]:
    """(status, failure_reason, page_count, token_estimate) of the stored row."""
    row = flow.db.attachment_row(attachment_id)
    assert row is not None
    return row["status"], row["failure_reason"], row["page_count"], row["token_estimate"]


def _derived_bytes(flow: _Flow, attachment_id: uuid.UUID) -> Any:
    """The stored row's derived_bytes (GH-188 §12.4)."""
    row = flow.db.attachment_row(attachment_id)
    assert row is not None
    return row["derived_bytes"]


def _manifest(flow: _Flow, attachment_id: uuid.UUID) -> dict[str, Any]:
    manifest: dict[str, Any] = json.loads(
        (_derived(flow, attachment_id) / "manifest.json").read_text(encoding="utf-8")
    )
    return manifest


def _platform(flow: _Flow, monkeypatch: pytest.MonkeyPatch, **columns: Any) -> None:
    """Store the platform row with these columns; empty the cache so it is read."""
    flow.db.add_platform_settings(**columns)
    monkeypatch.setattr(scoped_settings, "_platform_cache", None)


async def _process(ap: ModuleType, flow: _Flow, attachment_id: uuid.UUID) -> Any:
    """process_attachment with the DEFAULT processor (the real conversion)."""
    return await ap.process_attachment(flow.db.pool, flow.root, attachment_id, ORG_ID)


def _alive(pid: int) -> bool:
    try:
        os.kill(pid, 0)
    except ProcessLookupError:
        return False
    except PermissionError:
        return True
    return True


def _record_text(record: logging.LogRecord) -> str:
    """The record's message plus its traceback text, if it carries one."""
    return logging.Formatter("%(name)s %(levelname)s %(message)s").format(record)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


async def test_attachment_conversion_flow_mixed_pdf_ends_ready_with_manifest_and_parts(
    ap: ModuleType, flow: _Flow
) -> None:
    filename = f"Quartalsbericht Q3 {_EM_DASH} été.pdf"
    attachment_id = _store(flow, _mixed_pdf(), kind="pdf", filename=filename)

    result = await _process(ap, flow, attachment_id)

    status, reason, pages, estimate = _row(flow, attachment_id)
    derived = _derived(flow, attachment_id)
    manifest = _manifest(flow, attachment_id)
    parts = manifest["parts"]
    first = (derived / "part-0001.txt").read_text(encoding="utf-8")
    assert (result, status, reason, pages) == ("ready", "ready", None, 3)
    assert (manifest["version"], manifest["kind"], manifest["page_count"]) == (1, "pdf", 3)
    assert (estimate > 0, manifest["token_estimate"], sum(p["tokens"] for p in parts)) == (
        True,
        estimate,
        estimate,
    )
    assert sorted(os.listdir(derived)) == [
        "manifest.json",
        "part-0001.txt",
        "part-0002.jpg",
        "part-0003.txt",
    ]
    assert [(p["type"], p["page"], p["file"]) for p in parts] == [
        ("text", 1, "part-0001.txt"),
        ("image", 2, "part-0002.jpg"),
        ("text", 3, "part-0003.txt"),
    ]
    assert (
        first.startswith(f"[{filename} {_EM_DASH} page 1]\n"),
        "Quarterly revenue grew" in first,
        parts[1]["label"],
        parts[1]["media_type"],
    ) == (True, True, f"[{filename} {_EM_DASH} page 2]", "image/jpeg")


async def test_attachment_conversion_flow_pdf_over_the_page_limit_fails_too_many_pages(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch
) -> None:
    """#187's page limit, reported by the PDF converter: 3 pages over a stored limit
    of 2 end failed(too_many_pages); no artifacts are left, the original is kept."""
    _platform(flow, monkeypatch, max_pages_per_file=2)
    data = _mixed_pdf()
    attachment_id = _store(flow, data, kind="pdf", filename="long-report.pdf")

    result = await _process(ap, flow, attachment_id)

    assert (
        result,
        _row(flow, attachment_id),
        _gone(_derived(flow, attachment_id)),
        _original(flow, attachment_id).read_bytes() == data,
    ) == ("failed", ("failed", "too_many_pages", None, None), True, True)


async def test_attachment_conversion_flow_encrypted_pdf_fails_password_protected_without_a_worker(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    from admino.converters import runner

    marker = tmp_path / "worker-started"
    monkeypatch.setattr(
        runner,
        "WORKER_ARGV",
        (sys.executable, "-c", f"import pathlib; pathlib.Path({str(marker)!r}).touch()"),
    )
    attachment_id = _store(flow, _ENCRYPTED_PDF, kind="pdf", filename="locked.pdf")

    result = await _process(ap, flow, attachment_id)

    assert (
        result,
        _row(flow, attachment_id),
        marker.exists(),
        _gone(_derived(flow, attachment_id)),
    ) == ("failed", ("failed", "password_protected", None, None), False, True)


async def test_attachment_conversion_flow_scanned_page_rendered_at_the_stored_render_dpi(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The platform's render_dpi as stored (72, then 144) sets the rendered size of a
    144 x 216 pt scanned page: (144, 216) and (288, 432) px, each within 1 px; the
    JPEG on disk has the manifest's size."""
    from PIL import Image

    _platform(flow, monkeypatch, render_dpi=72)
    low = _store(flow, _scanned_pdf(), kind="pdf", filename="scan-low.pdf")
    await _process(ap, flow, low)
    flow.db.platform_settings[0]["render_dpi"] = 144
    monkeypatch.setattr(scoped_settings, "_platform_cache", None)
    high = _store(flow, _scanned_pdf(), kind="pdf", filename="scan-high.pdf")
    await _process(ap, flow, high)

    measured = []
    for attachment_id, (width, height) in ((low, (144, 216)), (high, (288, 432))):
        part = _manifest(flow, attachment_id)["parts"][0]
        with Image.open(_derived(flow, attachment_id) / part["file"]) as image:
            on_disk = image.size
        measured.append(
            (
                _row(flow, attachment_id)[:3],
                part["type"],
                abs(part["width"] - width) <= 1 and abs(part["height"] - height) <= 1,
                on_disk == (part["width"], part["height"]),
            )
        )
    assert measured == [(("ready", None, 1), "image", True, True)] * 2


# ---------------------------------------------------------------------------
# The other types
# ---------------------------------------------------------------------------


_TYPES: Final[dict[str, tuple[Callable[[], bytes], str, tuple[str, str], int | None]]] = {
    # kind: (fixture, file name, (part type, part file), exact estimate or None)
    "docx": (_docx, "budget.docx", ("text", "part-0001.txt"), None),
    "xlsx": (_xlsx, "figures.xlsx", ("text", "part-0001.txt"), None),
    "csv": (lambda: b"name;amount\nalpha;1\nbeta;2\n", "list.csv", ("text", "part-0001.txt"), None),
    "txt": (lambda: _TXT, "notes.txt", ("text", "part-0001.txt"), _text_tokens(_TXT_TEXT)),
    # 40 x 30 px: ceil(1200 / 750) = 2 tokens, no label.
    "png": (_png, "photo.png", ("image", "part-0001.png"), 2),
}


@pytest.mark.parametrize("kind", list(_TYPES))
async def test_attachment_conversion_flow_each_type_ends_ready_with_its_estimate(
    ap: ModuleType, flow: _Flow, kind: str
) -> None:
    fixture, filename, (part_type, part_file), exact = _TYPES[kind]
    attachment_id = _store(flow, fixture(), kind=kind, filename=filename)

    result = await _process(ap, flow, attachment_id)

    status, reason, pages, estimate = _row(flow, attachment_id)
    manifest = _manifest(flow, attachment_id)
    parts = manifest["parts"]
    assert (result, status, reason, pages, manifest["kind"], manifest["page_count"]) == (
        "ready",
        "ready",
        None,
        None,
        kind,
        None,
    )
    assert (
        estimate > 0,
        manifest["token_estimate"],
        sum(p["tokens"] for p in parts),
        estimate if exact is None else exact,
    ) == (True, estimate, estimate, estimate)
    assert (
        [(p["type"], p["file"]) for p in parts],
        sorted(os.listdir(_derived(flow, attachment_id))),
    ) == ([(part_type, part_file)], sorted(["manifest.json", part_file]))


# ---------------------------------------------------------------------------
# Derived bytes, the storage quota and the derived cap (GH-188 §12.4)
# ---------------------------------------------------------------------------


async def test_attachment_conversion_flow_ready_stores_the_derived_bytes_on_disk(
    ap: ModuleType, flow: _Flow
) -> None:
    """A real mixed PDF (two text parts, one JPEG, the manifest): derived_bytes is the
    sum of the sizes of the files in <id>.d."""
    attachment_id = _store(flow, _mixed_pdf(), kind="pdf", filename="mixed.pdf")

    result = await _process(ap, flow, attachment_id)

    derived = _derived(flow, attachment_id)
    on_disk = sum(entry.stat().st_size for entry in derived.iterdir())
    assert (result, len(os.listdir(derived)), on_disk > 0) == ("ready", 4, True)
    assert _derived_bytes(flow, attachment_id) == on_disk


async def test_attachment_conversion_flow_no_room_for_derived_files_fails_storage_quota(
    ap: ModuleType, flow: _Flow
) -> None:
    """The org's quota holds the original exactly, so its derived files don't fit:
    failed(storage_quota_exceeded), no <id>.d, no estimate or derived bytes stored, the
    original kept."""
    flow.db.add_org(ORG_ID, storage_quota_bytes=len(_TXT))
    attachment_id = _store(flow, _TXT, kind="txt", filename="notes.txt")

    result = await _process(ap, flow, attachment_id)

    assert (
        result,
        _row(flow, attachment_id),
        _derived_bytes(flow, attachment_id),
        _gone(_derived(flow, attachment_id)),
        _original(flow, attachment_id).read_bytes() == _TXT,
    ) == ("failed", ("failed", "storage_quota_exceeded", None, None), None, True, True)


async def test_attachment_conversion_flow_derived_files_over_the_cap_fail_output_too_large(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch
) -> None:
    """MAX_DERIVED_BYTES lowered to 16 in this (server) process: the child converts with
    its own cap, the runner measures the part and the manifest over 16 bytes:
    failed(output_too_large), no <id>.d, the original kept."""
    from admino.converters import common

    monkeypatch.setattr(common, "MAX_DERIVED_BYTES", 16)
    attachment_id = _store(flow, _TXT, kind="txt", filename="notes.txt")

    result = await _process(ap, flow, attachment_id)

    assert (
        result,
        _row(flow, attachment_id),
        _derived_bytes(flow, attachment_id),
        _gone(_derived(flow, attachment_id)),
        _original(flow, attachment_id).read_bytes() == _TXT,
    ) == ("failed", ("failed", "output_too_large", None, None), None, True, True)


# ---------------------------------------------------------------------------
# The worker process: timeout, the pool, logs
# ---------------------------------------------------------------------------


async def test_attachment_conversion_flow_timeout_kills_the_worker_and_fails(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A worker still running after CONVERSION_TIMEOUT_S (2 s here; it sleeps 60 s) is
    killed: its pid is gone right after, the file is failed(conversion_timeout), no
    artifacts are left and the original is kept."""
    from admino.converters import runner

    pid_file = tmp_path / "worker.pid"
    sleeper = (
        f"import os, pathlib, time; pathlib.Path({str(pid_file)!r})"
        ".write_text(str(os.getpid())); time.sleep(60)"
    )
    monkeypatch.setattr(runner, "WORKER_ARGV", (sys.executable, "-c", sleeper))
    monkeypatch.setattr(runner, "CONVERSION_TIMEOUT_S", 2.0)
    attachment_id = _store(flow, _TXT, kind="txt", filename="notes.txt")
    started = time.monotonic()
    pid: int | None = None
    try:
        result = await _process(ap, flow, attachment_id)
        elapsed = time.monotonic() - started
        pid = int(pid_file.read_text(encoding="utf-8")) if pid_file.exists() else None
        alive = pid is not None and _alive(pid)
    finally:
        if pid is not None and _alive(pid):
            os.kill(pid, signal.SIGKILL)

    assert (
        result,
        _row(flow, attachment_id),
        _gone(_derived(flow, attachment_id)),
        _original(flow, attachment_id).read_bytes() == _TXT,
    ) == ("failed", ("failed", "conversion_timeout", None, None), True, True)
    assert (pid is not None, alive, elapsed < 30.0) == (True, False, True)


async def test_attachment_conversion_flow_worker_flooding_stdout_is_killed_and_fails(
    ap: ModuleType,
    flow: _Flow,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """GH-281 Decision 1: a worker that writes 1 MiB to stdout (a marker, no newline) and
    then sleeps 60 s holding it open is killed as soon as more than runner.MAX_RESULT_BYTES
    have arrived: failed(processing_error) well before CONVERSION_TIMEOUT_S (8 s here),
    its pid gone right after (killed and reaped), no artifacts and no derived bytes, the
    original kept; no log record holds the flood."""
    from admino.converters import runner

    pid_file = tmp_path / "worker.pid"
    marker = "FLOOD-a41c"
    flooder = (
        f"import os, pathlib, sys, time; sys.stdin.buffer.read(); pathlib.Path({str(pid_file)!r})"
        f".write_text(str(os.getpid())); sys.stdout.buffer.write(b{marker!r} * 104858); "
        "sys.stdout.buffer.flush(); time.sleep(60)"
    )
    monkeypatch.setattr(runner, "WORKER_ARGV", (sys.executable, "-c", flooder))
    monkeypatch.setattr(runner, "CONVERSION_TIMEOUT_S", 8.0)
    caplog.set_level(logging.DEBUG)
    attachment_id = _store(flow, _TXT, kind="txt", filename="notes.txt")
    started = time.monotonic()
    pid: int | None = None
    try:
        result = await _process(ap, flow, attachment_id)
        elapsed = time.monotonic() - started
        pid = int(pid_file.read_text(encoding="utf-8")) if pid_file.exists() else None
        alive = pid is not None and _alive(pid)
    finally:
        if pid is not None and _alive(pid):
            os.kill(pid, signal.SIGKILL)

    assert (
        result,
        _row(flow, attachment_id),
        _derived_bytes(flow, attachment_id),
        _gone(_derived(flow, attachment_id)),
        _original(flow, attachment_id).read_bytes() == _TXT,
    ) == ("failed", ("failed", "processing_error", None, None), None, True, True)
    assert (pid is not None, alive, elapsed < 4.0, marker in caplog.text) == (
        True,
        False,
        True,
        False,
    )


async def test_attachment_conversion_flow_pool_converts_to_the_same_ready_outcome(
    ap: ModuleType, flow: _Flow
) -> None:
    attachment_id = _store(flow, _TXT, kind="txt", filename="notes.txt")
    processing = ap.ProcessingPool()

    processing.submit(flow.db.pool, flow.root, attachment_id, ORG_ID)
    await asyncio.wait_for(processing.join(), _JOIN_S)

    expected = _text_tokens(_TXT_TEXT)
    assert (
        _row(flow, attachment_id),
        _manifest(flow, attachment_id)["token_estimate"],
        sorted(os.listdir(_derived(flow, attachment_id))),
        (_derived(flow, attachment_id) / "part-0001.txt").read_text(encoding="utf-8"),
    ) == (("ready", None, None, expected), expected, ["manifest.json", "part-0001.txt"], _TXT_TEXT)


async def test_attachment_conversion_flow_logs_hold_no_file_name_path_or_text(
    ap: ModuleType, flow: _Flow, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """A ready TXT, an encrypted PDF (failed before any worker) and a worker that
    echoes its job to stdout and stderr and crashes (processing_error): nothing in
    the formatted output or the raw records (messages and tracebacks) names the file,
    a path under tmp_path or the file's text."""
    from admino.converters import runner

    name = "payroll-leak-name-7f3a"
    secret = "SECRET-CONTENT-91c2"
    ready_id = _store(flow, f"Salaries {secret}\n".encode(), kind="txt", filename=f"{name}.txt")
    locked_id = _store(flow, _ENCRYPTED_PDF, kind="pdf", filename=f"{name}-locked.pdf")
    crash_id = _store(flow, f"Bonus {secret}\n".encode(), kind="txt", filename=f"{name}-x.txt")

    with configured_logging("DEBUG", "json") as logs:
        logging.getLogger("admino.attachment_processing").debug("capture-check")
        outcomes = [await _process(ap, flow, ready_id), await _process(ap, flow, locked_id)]
        monkeypatch.setattr(runner, "WORKER_ARGV", _ECHO_AND_CRASH)
        outcomes.append(await _process(ap, flow, crash_id))

    seen = logs.text + "\n" + "\n".join(_record_text(record) for record in logs.records)
    markers = (name, secret, "Salaries", str(tmp_path), str(tmp_path.resolve()))
    assert (
        "capture-check" in logs.text,
        outcomes,
        [_row(flow, attachment_id)[1] for attachment_id in (ready_id, locked_id, crash_id)],
    ) == (True, ["ready", "failed", "failed"], [None, "password_protected", "processing_error"])
    assert {marker: marker in seen for marker in markers} == dict.fromkeys(markers, False)
