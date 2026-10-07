"""Tests for ``admino.converters.word.convert_docx`` (GH-188 contract section 4.2).

DOCX -> Markdown-like text (issue: "DOCX (python-docx): Markdown-like text with
headings, lists and tables"; Decision 5). Fixtures are built with python-docx
in ``tmp_path``; the converter is called in-process with a real ``PartWriter``
and the part file is read back. What these tests pin down:

- Paragraphs: ``clean_text(paragraph.text)``; empty (also whitespace-only)
  paragraphs skipped; line breaks kept, trailing spaces/tabs per line removed,
  tabs inside a line kept; C1 controls, DEL and U+FEFF removed.
- Styles: ``Title`` -> ``# ``; ``Heading N`` -> N ``#`` for 1..6, ``######``
  for 7..9; ``List Bullet``/``List Bullet N`` -> ``- `` items and ``List
  Number``/``List Number N`` -> ``1. `` items, indented by two spaces per level
  above 1; a paragraph with ``w:numPr`` (any other style) -> a ``- `` item at
  ``w:ilvl`` + 1 (missing ``w:ilvl`` -> level 1); the style wins over
  numbering (a numbered heading stays a heading); ``List Paragraph`` or
  ``List Continue`` without numbering stay prose; a multi-line item keeps its
  newlines.
- Joining: blocks by a blank line, two consecutive list items (any kinds,
  any levels) by one newline.
- Tables: Markdown with the first row as header, cells through
  ``clean_cell`` (pipe escaped, newlines and tabs -> one space), empty rows
  dropped (they don't count toward the row cap), trailing empty columns
  trimmed, caps with the exact notes (monkeypatched small
  ``MAX_TABLE_ROWS``/``MAX_TABLE_COLUMNS``/``MAX_CELL_CHARS``; no note at the
  caps); paragraphs and tables keep body order; an all-empty table adds no
  block (resolved gap: like an empty paragraph).
- Output: one text part ``part-0001.txt``, page None, tokens ==
  ``estimate_text_tokens`` of the exact text; returns None; an empty document
  writes no part.
- Failures: a broken ``word/document.xml``, a missing main part, a non-ZIP
  and an XLSX recorded as docx -> ``corrupted_file`` (message = the code);
  the zip-bomb guard runs first (``archive_too_large`` with a monkeypatched
  small entry limit, before any entry is read).
- XML entities in ``document.xml`` are never resolved: a ``file://`` external
  entity's secret never reaches the output; an entity-expansion bomb doesn't
  expand (``corrupted_file`` or a small text).

The converter modules are imported inside the helpers, so this file collects
before they exist.
"""

from __future__ import annotations

import zipfile
from typing import TYPE_CHECKING, Any

import docx
import pytest

if TYPE_CHECKING:
    from collections.abc import Callable, Mapping, Sequence
    from pathlib import Path

    from docx.document import Document as DocxDocument
    from docx.text.paragraph import Paragraph

_FILENAME = "Quartalsbericht Q3.docx"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_SECRET = "XXE-SECRET-MARKER-188"
_ELLIPSIS = chr(0x2026)


class _ReadEntryError(BaseException):
    """Raised when an archive entry is read where none may be (not an Exception,
    so the converter's error mapping can't swallow it)."""


def _options() -> Any:
    from admino.converters.common import ConversionOptions

    return ConversionOptions(filename=_FILENAME, render_dpi=150, max_pages=100)


def _save(tmp_path: Path, document: DocxDocument, name: str = "upload.docx") -> Path:
    path = tmp_path / name
    document.save(str(path))
    return path


def _run(tmp_path: Path, path: Path) -> tuple[Any, Any, Path]:
    """Convert ``path``; returns (return value, writer, out_dir)."""
    from admino.converters import common, word

    out_dir = tmp_path / "out"
    out_dir.mkdir()
    writer = common.PartWriter(out_dir)
    result = word.convert_docx(path, writer, _options())
    return result, writer, out_dir


def _text(tmp_path: Path, document: DocxDocument) -> str:
    """The single text part the document converts to."""
    result, writer, out_dir = _run(tmp_path, _save(tmp_path, document))
    assert (result, [part.file for part in writer.parts]) == (None, ["part-0001.txt"])
    return (out_dir / "part-0001.txt").read_text(encoding="utf-8")


def _refusal(tmp_path: Path, path: Path) -> tuple[str, str, list[str]]:
    """(reason, message, files left in out_dir) for a refused document."""
    from admino.converters.common import ConversionError

    with pytest.raises(ConversionError) as caught:
        _run(tmp_path, path)
    leftovers = sorted(entry.name for entry in (tmp_path / "out").iterdir())
    return caught.value.reason, str(caught.value), leftovers


def _numbered(paragraph: Paragraph, ilvl: int | None) -> Paragraph:
    """Give the paragraph direct numbering properties (``w:pPr/w:numPr``)."""
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    if ilvl is not None:
        num_pr.get_or_add_ilvl().val = ilvl
    num_pr.get_or_add_numId().val = 1
    return paragraph


def _table(document: DocxDocument, rows: Sequence[Sequence[str]]) -> Any:
    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    return table


def _rewrite(source: Path, target: Path, replace: Mapping[str, bytes | None]) -> Path:
    """Copy a DOCX, replacing (bytes) or dropping (None) the named entries."""
    with (
        zipfile.ZipFile(source) as original,
        zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as copy,
    ):
        for info in original.infolist():
            if info.filename not in replace:
                copy.writestr(info.filename, original.read(info.filename))
            elif (data := replace[info.filename]) is not None:
                copy.writestr(info.filename, data)
    return target


def _base_docx(tmp_path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("Hello")
    return _save(tmp_path, document, "base.docx")


# --- headings and paragraphs ----------------------------------------------------------


def test_word_title_and_headings_become_markdown_headings(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("Annual report", level=0)
    for level in range(1, 10):
        document.add_heading(f"Level {level}", level=level)

    expected = "\n\n".join(
        [
            "# Annual report",
            "# Level 1",
            "## Level 2",
            "### Level 3",
            "#### Level 4",
            "##### Level 5",
            "###### Level 6",
            "###### Level 7",
            "###### Level 8",
            "###### Level 9",
        ]
    )
    assert _text(tmp_path, document) == expected


def test_word_paragraphs_join_with_blank_line_and_empty_ones_skipped(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("")
    document.add_paragraph(" \t ")
    document.add_paragraph("Second paragraph.")

    assert _text(tmp_path, document) == "First paragraph.\n\nSecond paragraph."


def test_word_line_breaks_kept_and_line_ends_trimmed(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Line one ")
    paragraph.add_run().add_break()
    paragraph.add_run("Line\ttwo\t")

    assert _text(tmp_path, document) == "Line one\nLine\ttwo"


def test_word_control_characters_removed(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading(f"{chr(0xFEFF)}Budget", level=1)
    document.add_paragraph(f"a{chr(0x7F)}b{chr(0x85)}c{chr(0x9F)}d{chr(0xFEFF)}e")

    assert _text(tmp_path, document) == "# Budget\n\nabcde"


# --- lists ----------------------------------------------------------------------------


def test_word_list_styles_become_items_indented_by_level(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Apples", style="List Bullet")
    document.add_paragraph("Green", style="List Bullet 2")
    document.add_paragraph("Granny Smith", style="List Bullet 3")
    document.add_paragraph("Step one", style="List Number")
    document.add_paragraph("Sub step", style="List Number 2")
    document.add_paragraph("Detail", style="List Number 3")

    assert _text(tmp_path, document) == (
        "- Apples\n  - Green\n    - Granny Smith\n1. Step one\n  1. Sub step\n    1. Detail"
    )


def test_word_numbered_paragraphs_become_bullets_at_ilvl(tmp_path: Path) -> None:
    document = docx.Document()
    _numbered(document.add_paragraph("Word list item", style="List Paragraph"), ilvl=0)
    _numbered(document.add_paragraph("Nested item"), ilvl=2)
    _numbered(document.add_paragraph("No level given"), ilvl=None)

    assert _text(tmp_path, document) == "- Word list item\n    - Nested item\n- No level given"


def test_word_style_wins_over_numbering_and_list_lookalikes_stay_prose(tmp_path: Path) -> None:
    document = docx.Document()
    _numbered(document.add_heading("Numbered heading", level=2), ilvl=0)
    document.add_paragraph("Indented prose", style="List Paragraph")
    document.add_paragraph("Continued prose", style="List Continue")

    assert _text(tmp_path, document) == "## Numbered heading\n\nIndented prose\n\nContinued prose"


def test_word_consecutive_items_join_with_one_newline_other_blocks_with_blank_line(
    tmp_path: Path,
) -> None:
    document = docx.Document()
    document.add_paragraph("Intro")
    document.add_paragraph("one", style="List Bullet")
    document.add_paragraph("two", style="List Number")
    _numbered(document.add_paragraph("three"), ilvl=1)
    document.add_paragraph("Middle")
    document.add_paragraph("four", style="List Bullet")
    document.add_heading("Next", level=1)
    document.add_paragraph("five", style="List Bullet")
    _table(document, [["h"], ["v"]])

    assert _text(tmp_path, document) == (
        "Intro\n\n- one\n1. two\n  - three\n\nMiddle\n\n- four\n\n"
        "# Next\n\n- five\n\n| h |\n| --- |\n| v |"
    )


def test_word_multi_line_item_keeps_its_newlines(tmp_path: Path) -> None:
    document = docx.Document()
    item = document.add_paragraph("first line", style="List Bullet")
    item.add_run().add_break()
    item.add_run("second line")
    document.add_paragraph("next", style="List Bullet")

    assert _text(tmp_path, document) == "- first line\nsecond line\n- next"


# --- tables ---------------------------------------------------------------------------


def test_word_table_becomes_markdown_with_header_row_and_trailing_empty_column_trimmed(
    tmp_path: Path,
) -> None:
    document = docx.Document()
    _table(document, [["Name", "Amount", ""], ["Rent", "1200", ""], ["Food", "450", ""]])

    assert (
        _text(tmp_path, document)
        == "| Name | Amount |\n| --- | --- |\n| Rent | 1200 |\n| Food | 450 |"
    )


def test_word_table_cells_cleaned(tmp_path: Path) -> None:
    document = docx.Document()
    table = _table(document, [["Key", "Value"], ["a|b", "first"]])
    table.cell(1, 1).add_paragraph("second\tpart")
    table.cell(1, 0).paragraphs[0].add_run(chr(0x85))

    assert (
        _text(tmp_path, document) == "| Key | Value |\n| --- | --- |\n| a\\|b | first second part |"
    )


def test_word_table_empty_rows_dropped(tmp_path: Path) -> None:
    document = docx.Document()
    _table(document, [["H1", "H2"], ["", ""], ["a", "b"], [" ", ""]])

    assert _text(tmp_path, document) == "| H1 | H2 |\n| --- | --- |\n| a | b |"


def test_word_table_over_caps_cut_with_notes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from admino.converters import common

    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 2)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    monkeypatch.setattr(common, "MAX_CELL_CHARS", 5)
    document = docx.Document()
    _table(
        document,
        [["Name", "Value", "Extra"], ["alphabet", "1", "x"], ["beta", "2", ""], ["gamma", "3", ""]],
    )

    assert _text(tmp_path, document) == (
        f"| Name | Value |\n| --- | --- |\n| alpha{_ELLIPSIS} | 1 |\n\n"
        "[Only the first 2 rows are included.]\n[Only the first 2 columns are included.]"
    )


def test_word_table_at_caps_has_no_notes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from admino.converters import common

    monkeypatch.setattr(common, "MAX_TABLE_ROWS", 2)
    monkeypatch.setattr(common, "MAX_TABLE_COLUMNS", 2)
    document = docx.Document()
    _table(document, [["h1", "h2", ""], ["", "", ""], ["a", "b", ""], ["", "", ""]])

    assert _text(tmp_path, document) == "| h1 | h2 |\n| --- | --- |\n| a | b |"


def test_word_paragraphs_and_tables_keep_body_order(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Before")
    _table(document, [["k1", "v1"], ["a", "b"]])
    document.add_paragraph("Between")
    _table(document, [["k2"], ["c"]])
    document.add_paragraph("After")

    assert _text(tmp_path, document) == (
        "Before\n\n| k1 | v1 |\n| --- | --- |\n| a | b |\n\n"
        "Between\n\n| k2 |\n| --- |\n| c |\n\nAfter"
    )


def test_word_all_empty_table_adds_no_block(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Before")
    _table(document, [["", ""], [" ", ""]])
    document.add_paragraph("After")

    assert _text(tmp_path, document) == "Before\n\nAfter"


# --- output part ----------------------------------------------------------------------


def test_word_writes_one_text_part_with_page_none_and_its_token_estimate(tmp_path: Path) -> None:
    from admino.converters.common import TextPart
    from admino.tokens import estimate_text_tokens

    document = docx.Document()
    document.add_heading("Budget 2026", level=1)
    document.add_paragraph("Rent: 1200 CHF per month.")
    expected = "# Budget 2026\n\nRent: 1200 CHF per month."

    result, writer, out_dir = _run(tmp_path, _save(tmp_path, document))

    tokens = estimate_text_tokens(expected)
    assert (
        result,
        writer.parts,
        writer.token_estimate,
        (out_dir / "part-0001.txt").read_bytes(),
    ) == (
        None,
        [TextPart(type="text", file="part-0001.txt", page=None, tokens=tokens)],
        tokens,
        expected.encode("utf-8"),
    )


def _blank(document: DocxDocument) -> None:
    return None


def _only_empty_blocks(document: DocxDocument) -> None:
    document.add_paragraph("")
    document.add_paragraph("  ")
    _table(document, [["", ""], ["", " "]])


@pytest.mark.parametrize("fill", [_blank, _only_empty_blocks], ids=["blank", "only-empty-blocks"])
def test_word_empty_document_writes_no_part(
    tmp_path: Path, fill: Callable[[DocxDocument], None]
) -> None:
    document = docx.Document()
    fill(document)

    result, writer, out_dir = _run(tmp_path, _save(tmp_path, document))

    assert (result, writer.parts, sorted(out_dir.iterdir())) == (None, [], [])


# --- failures -------------------------------------------------------------------------


def _broken_document_xml(tmp_path: Path) -> Path:
    return _rewrite(
        _base_docx(tmp_path), tmp_path / "broken.docx", {"word/document.xml": b"<w:document"}
    )


def _missing_main_part(tmp_path: Path) -> Path:
    return _rewrite(_base_docx(tmp_path), tmp_path / "missing.docx", {"word/document.xml": None})


def _not_a_zip(tmp_path: Path) -> Path:
    path = tmp_path / "plain.docx"
    path.write_bytes(b"PK\x03\x04 this is not an archive")
    return path


def _workbook_as_docx(tmp_path: Path) -> Path:
    import openpyxl

    path = tmp_path / "book.docx"
    openpyxl.Workbook().save(path)
    return path


@pytest.mark.parametrize(
    "build",
    [_broken_document_xml, _missing_main_part, _not_a_zip, _workbook_as_docx],
    ids=["broken-document-xml", "missing-main-part", "not-a-zip", "xlsx-recorded-as-docx"],
)
def test_word_unreadable_document_is_corrupted_file(
    tmp_path: Path, build: Callable[[Path], Path]
) -> None:
    path = build(tmp_path)

    assert _refusal(tmp_path, path) == ("corrupted_file", "corrupted_file", [])


def test_word_zip_bomb_guard_runs_before_any_entry_is_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from admino.converters import common

    path = _base_docx(tmp_path)  # 17 entries
    monkeypatch.setattr(common, "MAX_OOXML_ENTRIES", 3)
    reads: list[str] = []
    real_open = zipfile.ZipFile.open

    def guarded_open(
        self: zipfile.ZipFile, name: Any, mode: str = "r", *args: Any, **kwargs: Any
    ) -> Any:
        if mode == "r":
            reads.append(str(name))
            raise _ReadEntryError(name)
        return real_open(self, name, mode, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", guarded_open)

    assert (_refusal(tmp_path, path), reads) == (("archive_too_large", "archive_too_large", []), [])


# --- XML entities never resolved ------------------------------------------------------


def _document_xml(doctype: str, text: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        f"{doctype}\n"
        f'<w:document xmlns:w="{_W_NS}"><w:body>'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:body></w:document>"
    ).encode()


def _outcome_and_written(tmp_path: Path, path: Path) -> tuple[str | None, str]:
    """(refusal reason or None, everything written to out_dir as text)."""
    from admino.converters.common import ConversionError

    try:
        _run(tmp_path, path)
    except ConversionError as exc:
        reason: str | None = exc.reason
    else:
        reason = None
    written = "".join(
        entry.read_text(encoding="utf-8", errors="replace")
        for entry in sorted((tmp_path / "out").rglob("*"))
        if entry.is_file()
    )
    return reason, written


def test_word_external_entity_never_resolved(tmp_path: Path) -> None:
    secret = tmp_path / "secret.txt"
    secret.write_text(_SECRET, encoding="utf-8")
    doctype = f'<!DOCTYPE w:document [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>'
    path = _rewrite(
        _base_docx(tmp_path),
        tmp_path / "xxe.docx",
        {"word/document.xml": _document_xml(doctype, "before &xxe; after")},
    )

    reason, written = _outcome_and_written(tmp_path, path)

    assert (reason in (None, "corrupted_file"), _SECRET in written) == (True, False)


def test_word_entity_expansion_bomb_not_expanded(tmp_path: Path) -> None:
    entities = ['<!ENTITY e0 "laughlaugh">']
    entities += [f'<!ENTITY e{n} "{f"&e{n - 1};" * 10}">' for n in range(1, 9)]
    doctype = f"<!DOCTYPE w:document [{''.join(entities)}]>"
    path = _rewrite(
        _base_docx(tmp_path),
        tmp_path / "lol.docx",
        {"word/document.xml": _document_xml(doctype, "x&e8;y")},
    )

    reason, written = _outcome_and_written(tmp_path, path)

    assert (reason in (None, "corrupted_file"), len(written) < 1000) == (True, True)
